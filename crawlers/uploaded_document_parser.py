# -*- coding: utf-8 -*-
"""Đọc tài liệu người dùng tải lên và chuyển thành dữ liệu tuyển sinh chuẩn."""

from html import escape
import ipaddress
import os
import re
import socket
from typing import Dict, Tuple
from urllib.parse import unquote, urljoin, urlparse

import requests
from werkzeug.utils import secure_filename
from core.models import (
    CrawlBundle,
    MethodConversionBundle,
    MethodEquivalenceRow,
    ScoreConversionRecord,
)
from crawlers.certificate_image_parser import rows_from_certificate_image
from crawlers.cutoff_image_parser import records_from_cutoff_image
from crawlers.dean_extractor import DeanAdmissionExtractor
from crawlers.table_image_parser import (
    equivalence_from_tokens,
    ocr_tokens,
    scalar_conversion_from_tokens,
)
from parsers.docx_parser import DocxAdmissionParser
from parsers.excel_parser import ExcelAdmissionParser
from parsers.pdf_parser import PdfAdmissionParser


_IMAGE_EXTS = {".jpg", ".jpeg", ".png", ".webp"}
_ALLOWED_EXTS = {".xlsx", ".xls", ".csv", ".pdf", ".docx", ".html", ".htm", *_IMAGE_EXTS}
_CONTENT_EXTS = {
    "application/pdf": ".pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": ".docx",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": ".xlsx",
    "application/vnd.ms-excel": ".xls",
    "text/csv": ".csv",
    "application/csv": ".csv",
    "text/comma-separated-values": ".csv",
    "image/jpeg": ".jpg",
    "image/png": ".png",
    "image/webp": ".webp",
    "text/html": ".html",
}


def _validate_public_url(url: str) -> None:
    parsed = urlparse(url)
    if parsed.scheme not in {"http", "https"} or not parsed.hostname:
        raise ValueError("chỉ hỗ trợ link http/https")
    if parsed.username or parsed.password:
        raise ValueError("link không được chứa thông tin đăng nhập")
    try:
        port = parsed.port
    except ValueError as exc:
        raise ValueError("cổng mạng không hợp lệ") from exc
    if port and port not in {80, 443}:
        raise ValueError("chỉ hỗ trợ cổng 80 hoặc 443")
    try:
        resolved_port = port or (443 if parsed.scheme == "https" else 80)
        addresses = {
            item[4][0] for item in socket.getaddrinfo(parsed.hostname, resolved_port)
        }
    except socket.gaierror as exc:
        raise ValueError("không phân giải được tên miền") from exc
    if not addresses:
        raise ValueError("không phân giải được tên miền")
    for address in addresses:
        ip = ipaddress.ip_address(address.split("%")[0])
        if not ip.is_global:
            raise ValueError("không cho phép địa chỉ mạng nội bộ")


def _cookie_challenge(body: bytes) -> Tuple[str, str]:
    """Một số cổng (daotao.neu.edu.vn) trả script đặt cookie rồi mới cho tải file."""
    head = body[:900].decode("utf-8", "replace")
    if "document.cookie" not in head or "location.reload" not in head:
        return "", ""
    match = re.search(
        r'document\.cookie\s*=\s*["\']([^="\'\s]+)=([^"\';]+)',
        head,
        re.I,
    )
    if not match:
        return "", ""
    return match.group(1), match.group(2)


def download_remote_document(
    url: str,
    upload_dir: str,
    prefix: str,
    max_bytes: int = 20 * 1024 * 1024,
) -> Dict[str, object]:
    """Tải tài liệu CDN an toàn, có kiểm tra redirect, định dạng và dung lượng."""
    current = (url or "").strip()
    response = None
    session = requests.Session()
    session.headers["User-Agent"] = "Mozilla/5.0 AdmissionDocumentImporter/1.0"
    verify = True
    for _ in range(6):
        _validate_public_url(current)
        try:
            response = session.get(
                current,
                timeout=(8, 30),
                allow_redirects=False,
                verify=verify,
            )
        except requests.exceptions.SSLError:
            verify = False
            response = session.get(
                current,
                timeout=(8, 30),
                allow_redirects=False,
                verify=False,
            )
        if response.status_code in {301, 302, 303, 307, 308}:
            target = response.headers.get("Location") or ""
            if not target:
                raise ValueError("link chuyển hướng không hợp lệ")
            current = urljoin(current, target)
            continue
        if response.status_code != 200:
            raise ValueError(f"máy chủ CDN trả về HTTP {response.status_code}")
        cookie_name, cookie_value = _cookie_challenge(response.content[:900])
        if cookie_name:
            host = urlparse(current).hostname or ""
            session.cookies.set(cookie_name, cookie_value, domain=host, path="/")
            continue
        break
    else:
        raise ValueError("link chuyển hướng quá nhiều lần")
    if response is None:
        raise ValueError("không tải được tài liệu")

    try:
        declared = int(response.headers.get("Content-Length") or 0)
    except (TypeError, ValueError):
        declared = 0
    if declared > max_bytes:
        response.close()
        raise ValueError("file vượt quá 20 MB")
    content_type = (response.headers.get("Content-Type") or "").split(";", 1)[0].lower()
    path_name = unquote(os.path.basename(urlparse(current).path))
    disposition = response.headers.get("Content-Disposition") or ""
    match = re.search(r"filename\*?=(?:UTF-8''|[\"']?)([^\"';]+)", disposition, re.I)
    original_name = unquote(match.group(1).strip()) if match else path_name
    ext = os.path.splitext(original_name)[1].lower()
    if ext not in _ALLOWED_EXTS:
        ext = _CONTENT_EXTS.get(content_type, "")
    if ext not in _ALLOWED_EXTS:
        response.close()
        raise ValueError("định dạng tài liệu CDN không được hỗ trợ")
    base = secure_filename(os.path.splitext(original_name)[0]) or "cdn_document"
    original_name = f"{base}{ext}"
    stored_name = f"{prefix}_{original_name}"
    path = os.path.join(upload_dir, stored_name)
    body = response.content or b""
    if len(body) > max_bytes:
        raise ValueError("file vượt quá 20 MB")
    if ext == ".pdf" and not body.startswith(b"%PDF"):
        raise ValueError("link không trả về file PDF")
    size = len(body)
    try:
        with open(path, "wb") as output:
            output.write(body)
    except Exception:
        if os.path.exists(path):
            os.unlink(path)
        raise
    return {
        "path": path,
        "filename": original_name,
        "stored_name": stored_name,
        "size": size,
        "source_url": current,
    }


def _source(filename: str, source_url: str) -> str:
    return f"Tài liệu tải lên: {filename} · {source_url}"


def _docx_html(path: str) -> str:
    import docx

    doc = docx.Document(path)
    chunks = [f"<p>{escape(p.text)}</p>" for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        chunks.append("<table>")
        for row in table.rows:
            chunks.append("<tr>")
            chunks.extend(f"<td>{escape(cell.text)}</td>" for cell in row.cells)
            chunks.append("</tr>")
        chunks.append("</table>")
    return "".join(chunks)


def _pdf_html(path: str, max_pages: int = 50) -> str:
    import pdfplumber

    chunks = []
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages[:max_pages]:
            text = page.extract_text() or ""
            chunks.extend(f"<p>{escape(line)}</p>" for line in text.splitlines() if line.strip())
            for table in page.extract_tables() or []:
                chunks.append("<table>")
                for row in table or []:
                    chunks.append("<tr>")
                    chunks.extend(f"<td>{escape(str(cell or ''))}</td>" for cell in row)
                    chunks.append("</tr>")
                chunks.append("</table>")
    return "".join(chunks)


def _dedupe_admissions(records):
    found = []
    seen = set()
    for row in records:
        key = (
            row.ma_nganh,
            row.ten_nganh,
            row.nam,
            row.phuong_thuc,
            row.to_hop,
            row.diem_chuan,
            row.diem_chuan_ptxt,
            row.chi_tieu,
        )
        if key not in seen:
            seen.add(key)
            found.append(row)
    return found


def parse_uploaded_document(
    path: str,
    filename: str,
    school_code: str,
    school_name: str,
    year: int,
    source_url: str,
    admission_method: str = "",
) -> Tuple[CrawlBundle, MethodConversionBundle]:
    """Phân tích một file đã lưu; không truy cập website bên ngoài."""
    crawl = CrawlBundle()
    method = MethodConversionBundle()
    ext = os.path.splitext(filename)[1].lower()
    source = _source(filename, source_url)

    if ext in {".xlsx", ".xls", ".csv"}:
        crawl.admissions.extend(
            ExcelAdmissionParser().parse(path, school_code, school_name, year)
        )
    elif ext in {".html", ".htm"}:
        with open(path, "r", encoding="utf-8", errors="replace") as handle:
            html = handle.read()
        admissions, conversions, regulations = DeanAdmissionExtractor().extract(
            html, school_code, school_name, source, year
        )
        crawl.admissions.extend(admissions)
        crawl.conversions.extend(conversions)
        crawl.regulations.extend(regulations)
    elif ext == ".pdf":
        crawl.admissions.extend(
            PdfAdmissionParser().parse(path, school_code, school_name, year, max_pages=50)
        )
        try:
            html = _pdf_html(path)
            admissions, conversions, regulations = DeanAdmissionExtractor().extract(
                html, school_code, school_name, source, year
            )
            crawl.admissions.extend(admissions)
            crawl.conversions.extend(conversions)
            crawl.regulations.extend(regulations)
        except Exception:
            pass
        if not crawl.admissions:
            from crawlers.official_site_crawler import OfficialSiteCrawler
            crawl.admissions.extend(
                OfficialSiteCrawler()._records_from_scanned_pdf(
                    path, school_code, school_name, year, source
                )
            )
    elif ext == ".docx":
        crawl.admissions.extend(
            DocxAdmissionParser().parse(path, school_code, school_name, year)
        )
        try:
            admissions, conversions, regulations = DeanAdmissionExtractor().extract(
                _docx_html(path), school_code, school_name, source, year
            )
            crawl.admissions.extend(admissions)
            crawl.conversions.extend(conversions)
            crawl.regulations.extend(regulations)
        except Exception:
            pass
    elif ext in _IMAGE_EXTS:
        crawl.admissions.extend(
            records_from_cutoff_image(path, school_code, school_name, year, source)
        )
        tokens = ocr_tokens(path)
        title, rows = equivalence_from_tokens(tokens, filename)
        if not rows:
            title, rows = scalar_conversion_from_tokens(tokens, filename)
        for values in rows:
            stt = values.pop("__stt", "")
            method.rows.append(
                MethodEquivalenceRow(
                    ma_truong=school_code,
                    ten_truong=school_name,
                    tieu_de_bang=title or filename,
                    stt=stt,
                    cot_gia_tri=values,
                    nam=year,
                    url_nguon=source_url,
                    nguon="Tài liệu tải lên",
                )
            )
        try:
            cert_rows = rows_from_certificate_image(path)
        except Exception:
            cert_rows = []
        for item in cert_rows:
            crawl.conversions.append(
                ScoreConversionRecord(
                    ma_truong=school_code,
                    ten_truong=school_name,
                    loai_bang=item["loai_bang"],
                    hang_muc=item["hang_muc"],
                    diem_quy_doi=item["diem_quy_doi"],
                    chi_tiet_hang=item.get("chi_tiet_hang") or "",
                    phuong_thuc=item.get("phuong_thuc") or "",
                    thang_diem=item.get("thang_diem") or "",
                    nam=year,
                    nguon=source,
                )
            )

    crawl.admissions = _dedupe_admissions(crawl.admissions)
    chosen_method = (admission_method or "").strip()
    for row in crawl.admissions:
        row.ma_truong = school_code
        row.ten_truong = school_name
        row.nguon = source
        if chosen_method:
            row.phuong_thuc = chosen_method
    for row in crawl.conversions:
        row.ma_truong = school_code
        row.ten_truong = school_name
        row.nguon = source
    for row in crawl.regulations:
        row.ma_truong = school_code
        row.ten_truong = school_name
        row.nguon = source
    return crawl, method
