# -*- coding: utf-8 -*-
"""
Module: crawlers.score_conversion_crawler
Mô tả: Cào trang Quy đổi điểm giữa các phương thức xét tuyển từ tuyensinh247.

Ví dụ:
  https://diemthi.tuyensinh247.com/quy-doi-diem/dai-hoc-kinh-te-quoc-dan-KHA.html

Thu thập:
  - Bảng tương đương THPT ↔ HSA ↔ TSA ↔ V-ACT ↔ SAT ↔ học bạ / kết hợp,...
  - Ghi chú / công thức / chênh lệch tổ hợp
  - Ảnh bảng quy đổi (khi trường chỉ đăng ảnh, VD: BKA)
  - Khoảng điểm đầu vào của công cụ quy đổi (HSA 85-150,...)
"""

import os
import re
import json
from datetime import datetime
from typing import List, Dict, Optional, Any, Tuple
from html import unescape

import requests
from bs4 import BeautifulSoup, Tag
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from core.models import (
    MethodEquivalenceRow,
    MethodConversionNote,
    MethodConversionImage,
    MethodRangeHint,
    MethodConversionBundle,
)
from core.normalizer import clean_text, extract_year_from_text, strip_accents
from crawlers.school_directory import SchoolDirectory


class ScoreConversionCrawler:
    """Crawler trang /quy-doi-diem/{slug}.html."""

    BASE_URL = "https://diemthi.tuyensinh247.com"
    HEADERS = {
        "User-Agent": (
            "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) "
            "AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36"
        )
    }

    METHOD_RANGE_KEYS = [
        "HSA", "TSA", "SAT", "ACT", "V-ACT", "VACT", "ĐGNL", "DGNL",
        "ĐGTD", "DGTD", "IELTS", "TOEFL", "TOEIC", "XTTN",
    ]

    def __init__(self, cache_dir: str = "data/cache"):
        self.cache_dir = cache_dir
        os.makedirs(self.cache_dir, exist_ok=True)
        self.directory = SchoolDirectory(cache_dir=cache_dir)
        self.directory.load()

    def resolve_school(self, code_or_name: str) -> Optional[Dict[str, Any]]:
        code = (code_or_name or "").strip().upper()
        # Alias phổ biến
        aliases = {"NEU": "KHA", "FTU": "NTH", "HUST": "BKA", "UET": "QHI"}
        code = aliases.get(code, code)
        info = self.directory.schools.get(code)
        if info:
            return info
        # Tìm theo tên
        key = (code_or_name or "").strip().lower()
        for c, inf in self.directory.schools.items():
            if key in (inf.get("name") or "").lower() or key in (inf.get("slug") or "").lower():
                return inf
        return None

    def crawl_school(
        self,
        school_code_or_name: str,
    ) -> Tuple[MethodConversionBundle, Dict[str, Any]]:
        """Cào quy đổi điểm của 1 trường. Trả về (bundle, meta)."""
        bundle = MethodConversionBundle()
        info = self.resolve_school(school_code_or_name)
        if not info:
            meta = {
                "ok": False,
                "code": school_code_or_name,
                "error": f"Không tìm thấy trường '{school_code_or_name}' trong danh bạ.",
            }
            return bundle, meta

        code = info["code"]
        name = info.get("name") or code
        slug = info.get("slug") or ""
        url = f"{self.BASE_URL}/quy-doi-diem/{slug}.html"
        meta = {
            "ok": False,
            "code": code,
            "name": name,
            "slug": slug,
            "url": url,
            "tables": 0,
            "notes": 0,
            "images": 0,
            "ranges": 0,
            "error": "",
        }

        try:
            res = requests.get(url, headers=self.HEADERS, timeout=20)
            if res.status_code != 200:
                meta["error"] = f"HTTP {res.status_code}"
                return bundle, meta

            soup = BeautifulSoup(res.text, "html.parser")
            year = extract_year_from_text(soup.get_text(" ", strip=True)[:2000]) or 2026

            rows = self._extract_tables(soup, code, name, year, url)
            notes = self._extract_notes(soup, code, name, year, url)
            images = self._extract_images(soup, code, name, year, url)
            ranges = self._extract_method_ranges(soup, code, name, year, url)

            # Fallback: nếu không có bảng HTML, thử bóc từ RSC payload
            if not rows:
                rows = self._extract_tables_from_rsc(res.text, code, name, year, url)

            bundle.rows.extend(rows)
            bundle.notes.extend(notes)
            bundle.images.extend(images)
            bundle.ranges.extend(ranges)

            meta.update({
                "ok": True,
                "tables": len({r.tieu_de_bang for r in rows}),
                "row_count": len(rows),
                "notes": len(notes),
                "images": len(images),
                "ranges": len(ranges),
            })
        except Exception as e:
            meta["error"] = str(e)

        bundle.school_results.append(meta)
        return bundle, meta

    def crawl_schools(self, school_codes: List[str]) -> MethodConversionBundle:
        """Cào nhiều trường, gộp kết quả."""
        combined = MethodConversionBundle()
        for code in school_codes:
            part, meta = self.crawl_school(code)
            combined.rows.extend(part.rows)
            combined.notes.extend(part.notes)
            combined.images.extend(part.images)
            combined.ranges.extend(part.ranges)
            combined.school_results.append(meta)
        return combined

    def _nearby_title(self, table: Tag) -> str:
        for tag in ["h2", "h3", "h4", "strong", "b", "p"]:
            prev = table.find_previous(tag)
            if prev:
                text = clean_text(prev.get_text(" ", strip=True))
                if text and len(text) < 220:
                    return text
        return "Bảng quy đổi điểm"

    def _extract_tables(
        self,
        soup: BeautifulSoup,
        code: str,
        name: str,
        year: Optional[int],
        url: str,
    ) -> List[MethodEquivalenceRow]:
        results: List[MethodEquivalenceRow] = []
        for table in soup.find_all("table"):
            rows = table.find_all("tr")
            if len(rows) < 2:
                continue
            headers = [clean_text(td.get_text(" ", strip=True)) for td in rows[0].find_all(["th", "td"])]
            if len(headers) < 2:
                continue

            # Chỉ lấy bảng có dấu hiệu quy đổi phương thức / điểm
            hjoin = strip_accents(" ".join(headers))
            if not any(
                k in hjoin
                for k in [
                    "thpt", "hsa", "tsa", "sat", "act", "v-act", "vact",
                    "dgnl", "dgtd", "hoc ba", "ket hop", "ielts", "toefl",
                    "diem", "quy doi", "xttn",
                ]
            ):
                continue

            title = self._nearby_title(table)
            for tr in rows[1:]:
                cells = [clean_text(td.get_text(" ", strip=True)) for td in tr.find_all(["th", "td"])]
                if not cells or all(not c for c in cells):
                    continue
                if any("tuyensinh247" in c.lower() for c in cells):
                    continue
                # Bỏ dòng header lặp
                if strip_accents(cells[0]) in ["tt", "stt"] and len(cells) > 1 and "thpt" in strip_accents(cells[1]):
                    continue

                col_map: Dict[str, str] = {}
                stt = ""
                for i, h in enumerate(headers):
                    if i >= len(cells):
                        break
                    hn = strip_accents(h)
                    if hn in ["tt", "stt"]:
                        stt = cells[i]
                        continue
                    if cells[i]:
                        col_map[h] = cells[i]

                if not col_map:
                    continue

                results.append(
                    MethodEquivalenceRow(
                        ma_truong=code,
                        ten_truong=name,
                        tieu_de_bang=title,
                        stt=stt,
                        cot_gia_tri=col_map,
                        nam=year,
                        url_nguon=url,
                    )
                )
        return results

    def _extract_tables_from_rsc(
        self,
        html: str,
        code: str,
        name: str,
        year: Optional[int],
        url: str,
    ) -> List[MethodEquivalenceRow]:
        """Thử bóc bảng từ payload Next.js nếu HTML tĩnh không có <table>."""
        results: List[MethodEquivalenceRow] = []
        # Tìm các đoạn HTML bị escape trong RSC
        fragments = re.findall(r"<table[\s\S]{20,8000}?</table>", html, flags=re.I)
        for frag in fragments:
            try:
                decoded = unescape(frag.encode("utf-8").decode("unicode_escape", errors="ignore"))
            except Exception:
                decoded = unescape(frag)
            soup = BeautifulSoup(decoded, "html.parser")
            results.extend(self._extract_tables(soup, code, name, year, url))
        return results

    def _extract_notes(
        self,
        soup: BeautifulSoup,
        code: str,
        name: str,
        year: Optional[int],
        url: str,
    ) -> List[MethodConversionNote]:
        notes: List[MethodConversionNote] = []
        seen = set()
        keywords = [
            r"quy\s*đổi",
            r"chênh\s*lệch",
            r"tổ\s*hợp",
            r"tương\s*đương",
            r"phương\s*thức",
            r"công\s*thức",
            r"ielts",
            r"chứng\s*chỉ",
            r"học\s*bạ",
            r"kết\s*hợp",
            r"hsa",
            r"tsa",
            r"v-act",
            r"sat",
        ]

        for el in soup.find_all(["h2", "h3", "h4", "strong", "p"]):
            title = clean_text(el.get_text(" ", strip=True))
            if not title or len(title) > 200:
                continue
            title_l = title.lower()
            if not any(re.search(k, title_l) for k in keywords):
                continue

            body_parts: List[str] = []
            if el.name == "p" and len(title) > 60:
                body = title
                short_title = title[:90] + ("..." if len(title) > 90 else "")
            else:
                short_title = title
                sib = el.find_next_sibling()
                steps = 0
                while sib is not None and steps < 5:
                    if getattr(sib, "name", None) in ["h2", "h3", "h4", "table"]:
                        break
                    txt = clean_text(sib.get_text(" ", strip=True)) if hasattr(sib, "get_text") else ""
                    if txt and len(txt) > 25:
                        body_parts.append(txt)
                    sib = sib.find_next_sibling()
                    steps += 1
                if not body_parts:
                    nxt = el.find_next("p")
                    if nxt:
                        body_parts.append(clean_text(nxt.get_text(" ", strip=True)))
                body = " ".join(body_parts).strip()

            if not body or len(body) < 40:
                continue
            if len(body) > 2500:
                body = body[:2500] + "..."

            key = (short_title[:50], body[:80])
            if key in seen:
                continue
            seen.add(key)
            notes.append(
                MethodConversionNote(
                    ma_truong=code,
                    ten_truong=name,
                    tieu_de=short_title,
                    noi_dung=body,
                    nam=year,
                    url_nguon=url,
                )
            )
        return notes

    def _extract_images(
        self,
        soup: BeautifulSoup,
        code: str,
        name: str,
        year: Optional[int],
        url: str,
    ) -> List[MethodConversionImage]:
        images: List[MethodConversionImage] = []
        seen = set()
        for img in soup.find_all("img"):
            src = img.get("src") or img.get("data-src") or ""
            if not src:
                continue
            src_l = src.lower()
            alt = clean_text(img.get("alt") or "")
            # Ảnh quy đổi / bảng điểm
            if not any(k in src_l or k in alt.lower() for k in ["quy-doi", "quy_doi", "quydoi", "convert", "bang"]):
                # vẫn lấy ảnh CDN picture năm tuyển sinh nếu nằm trong nội dung chính
                if "cdn.tuyensinh247.com/picture" not in src_l:
                    continue
                if not any(k in src_l for k in ["quy", "doi", "convert", "bang", "tuong-quan", "phan-vi"]):
                    continue
            if src in seen:
                continue
            seen.add(src)
            if src.startswith("//"):
                src = "https:" + src
            elif src.startswith("/"):
                src = self.BASE_URL + src
            images.append(
                MethodConversionImage(
                    ma_truong=code,
                    ten_truong=name,
                    url_anh=src,
                    mo_ta=alt or "Ảnh bảng quy đổi điểm",
                    nam=year,
                    url_nguon=url,
                )
            )
        return images

    def _extract_method_ranges(
        self,
        soup: BeautifulSoup,
        code: str,
        name: str,
        year: Optional[int],
        url: str,
    ) -> List[MethodRangeHint]:
        """Lấy khoảng điểm kiểu HSA 85-150 từ khối công cụ quy đổi."""
        ranges: List[MethodRangeHint] = []
        text = soup.get_text("\n", strip=True)
        lines = [ln.strip() for ln in text.split("\n") if ln.strip()]

        method_re = re.compile(
            r"^(HSA|SAT|TSA|V-?ACT|ACT|ĐGNL|DGNL|ĐGTD|DGTD|IELTS|TOEFL|XTTN)$",
            flags=re.I,
        )
        range_re = re.compile(
            r"^(\d{2,4}(?:[.,]\d+)?\s*[-–—]\s*\d{2,4}(?:[.,]\d+)?)$"
        )

        seen = set()
        # 1) METHOD và khoảng trên 2 dòng liền nhau
        for i, ln in enumerate(lines):
            m = method_re.match(ln)
            if not m or i + 1 >= len(lines):
                continue
            rm = range_re.match(lines[i + 1])
            if not rm:
                continue
            method = m.group(1).upper().replace("VACT", "V-ACT")
            khoang = re.sub(r"\s+", "", rm.group(1).replace("–", "-").replace("—", "-"))
            key = (method, khoang)
            if key in seen:
                continue
            seen.add(key)
            ranges.append(
                MethodRangeHint(
                    ma_truong=code,
                    ten_truong=name,
                    phuong_thuc=method,
                    khoang_diem=khoang,
                    nam=year,
                    url_nguon=url,
                )
            )

        # 2) METHOD và khoảng trên cùng dòng
        pattern = re.compile(
            r"\b(HSA|SAT|TSA|V-?ACT|ACT|ĐGNL|DGNL|ĐGTD|DGTD|IELTS|TOEFL|XTTN)\b"
            r"[^\n\d]{0,40}?"
            r"(\d{2,4}(?:[.,]\d+)?\s*[-–—]\s*\d{2,4}(?:[.,]\d+)?)",
            flags=re.I,
        )
        for m in pattern.finditer(text):
            method = m.group(1).upper().replace("VACT", "V-ACT")
            khoang = re.sub(r"\s+", "", m.group(2).replace("–", "-").replace("—", "-"))
            key = (method, khoang)
            if key in seen:
                continue
            seen.add(key)
            ranges.append(
                MethodRangeHint(
                    ma_truong=code,
                    ten_truong=name,
                    phuong_thuc=method,
                    khoang_diem=khoang,
                    nam=year,
                    url_nguon=url,
                )
            )
        return ranges

    def export_excel(
        self,
        bundle: MethodConversionBundle,
        output_path: str = "data/output/quy_doi_diem.xlsx",
    ) -> str:
        """Xuất Excel: Bang_Quy_Doi | Ghi_Chu | Anh | Khoang_Diem | Tom_Tat."""
        os.makedirs(os.path.dirname(output_path) or ".", exist_ok=True)
        wb = Workbook()

        header_fill = PatternFill(start_color="0D7377", end_color="0D7377", fill_type="solid")
        header_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        data_font = Font(name="Arial", size=10)
        thin = Border(
            left=Side(style="thin", color="D3D3D3"),
            right=Side(style="thin", color="D3D3D3"),
            top=Side(style="thin", color="D3D3D3"),
            bottom=Side(style="thin", color="D3D3D3"),
        )

        def style_header(ws, headers):
            for col, title in enumerate(headers, start=1):
                cell = ws.cell(row=1, column=col, value=title)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(ws.max_row, 2)}"

        # Sheet 1: bảng quy đổi (flatten cột động)
        ws1 = wb.active
        ws1.title = "Bang_Quy_Doi"
        all_cols = []
        for r in bundle.rows:
            for k in r.cot_gia_tri.keys():
                if k not in all_cols:
                    all_cols.append(k)
        headers1 = ["Mã trường", "Tên trường", "Tiêu đề bảng", "STT", "Năm", "URL nguồn"] + all_cols
        style_header(ws1, headers1)
        for i, r in enumerate(bundle.rows, start=2):
            vals = [r.ma_truong, r.ten_truong, r.tieu_de_bang, r.stt, r.nam or "", r.url_nguon]
            vals += [r.cot_gia_tri.get(c, "") for c in all_cols]
            for c, v in enumerate(vals, start=1):
                cell = ws1.cell(row=i, column=c, value=v)
                cell.font = data_font
                cell.border = thin
                cell.alignment = Alignment(wrap_text=True, vertical="center")
        if not bundle.rows:
            ws1.cell(row=2, column=1, value="Không có bảng HTML (có thể trường chỉ đăng ảnh — xem sheet Anh).")

        # Sheet 2: ghi chú / công thức
        ws2 = wb.create_sheet("Ghi_Chu_Cong_Thuc")
        headers2 = ["Mã trường", "Tên trường", "Tiêu đề", "Nội dung", "Năm", "URL nguồn"]
        style_header(ws2, headers2)
        for i, n in enumerate(bundle.notes, start=2):
            for c, v in enumerate([n.ma_truong, n.ten_truong, n.tieu_de, n.noi_dung, n.nam or "", n.url_nguon], start=1):
                cell = ws2.cell(row=i, column=c, value=v)
                cell.font = data_font
                cell.border = thin
                cell.alignment = Alignment(wrap_text=True, vertical="top")
            ws2.row_dimensions[i].height = min(90, 20 + len(n.noi_dung) // 90 * 12)
        ws2.column_dimensions["D"].width = 80

        # Sheet 3: ảnh
        ws3 = wb.create_sheet("Anh_Quy_Doi")
        headers3 = ["Mã trường", "Tên trường", "Mô tả", "URL ảnh", "Năm", "URL trang"]
        style_header(ws3, headers3)
        for i, img in enumerate(bundle.images, start=2):
            for c, v in enumerate([img.ma_truong, img.ten_truong, img.mo_ta, img.url_anh, img.nam or "", img.url_nguon], start=1):
                cell = ws3.cell(row=i, column=c, value=v)
                cell.font = data_font
                cell.border = thin

        # Sheet 4: khoảng điểm công cụ
        ws4 = wb.create_sheet("Khoang_Diem_Cong_Cu")
        headers4 = ["Mã trường", "Tên trường", "Phương thức", "Khoảng điểm", "Năm", "URL"]
        style_header(ws4, headers4)
        for i, rg in enumerate(bundle.ranges, start=2):
            for c, v in enumerate([rg.ma_truong, rg.ten_truong, rg.phuong_thuc, rg.khoang_diem, rg.nam or "", rg.url_nguon], start=1):
                cell = ws4.cell(row=i, column=c, value=v)
                cell.font = data_font
                cell.border = thin

        # Sheet 5: tóm tắt theo trường
        ws5 = wb.create_sheet("Tom_Tat_Theo_Truong")
        headers5 = ["Mã", "Tên", "OK", "Số dòng bảng", "Ghi chú", "Ảnh", "Khoảng điểm", "URL", "Lỗi"]
        style_header(ws5, headers5)
        for i, m in enumerate(bundle.school_results, start=2):
            vals = [
                m.get("code"), m.get("name"), "✓" if m.get("ok") else "✗",
                m.get("row_count", 0), m.get("notes", 0), m.get("images", 0),
                m.get("ranges", 0), m.get("url", ""), m.get("error", ""),
            ]
            for c, v in enumerate(vals, start=1):
                cell = ws5.cell(row=i, column=c, value=v)
                cell.font = data_font
                cell.border = thin

        for ws in [ws1, ws2, ws3, ws4, ws5]:
            for col in range(1, ws.max_column + 1):
                letter = get_column_letter(col)
                if ws.column_dimensions[letter].width is None or ws.column_dimensions[letter].width < 12:
                    ws.column_dimensions[letter].width = 16

        wb.save(output_path)
        return output_path

    def bundle_to_api_dict(self, bundle: MethodConversionBundle) -> Dict[str, Any]:
        """Serialize cho API/UI."""
        return {
            "rows": [r.to_dict() for r in bundle.rows],
            "notes": [n.to_dict() for n in bundle.notes],
            "images": [i.to_dict() for i in bundle.images],
            "ranges": [g.to_dict() for g in bundle.ranges],
            "school_results": bundle.school_results,
            "summary": {
                "schools": len(bundle.school_results),
                "rows": len(bundle.rows),
                "notes": len(bundle.notes),
                "images": len(bundle.images),
                "ranges": len(bundle.ranges),
            },
        }
