# -*- coding: utf-8 -*-
"""
Nhận diện 6 nhóm phương thức tuyển sinh trên website / đề án của trường.
"""

import json
import os
import re
import tempfile
from typing import Dict, List, Optional
from urllib.parse import urlparse

from bs4 import BeautifulSoup

from core.normalizer import clean_text, strip_accents
from crawlers.official_site_crawler import (
    OfficialSiteCrawler,
    _blob,
    _canon,
    _years_in,
    admission_portal_urls,
    filter_official_urls,
    normalize_site_url,
)


METHOD_CATALOG: List[Dict[str, str]] = [
    {
        "id": "thpt",
        "nhan": "THPT",
        "ten": "Xét tuyển bằng kết quả Kỳ thi tốt nghiệp THPT",
        "mo_ta": (
            "Sử dụng điểm số các bài thi/môn thi trong kỳ thi tốt nghiệp THPT quốc gia "
            "do Bộ Giáo dục và Đào tạo tổ chức theo các tổ hợp môn xét tuyển."
        ),
    },
    {
        "id": "hoc_ba",
        "nhan": "Học bạ",
        "ten": "Xét tuyển bằng học bạ THPT",
        "mo_ta": (
            "Dựa vào kết quả học tập (điểm trung bình các môn học) trong học bạ cấp THPT "
            "của các năm lớp 10, lớp 11 và lớp 12."
        ),
    },
    {
        "id": "dgnl",
        "nhan": "ĐGNL/ĐGTD",
        "ten": "Xét tuyển bằng kết quả Kỳ thi Đánh giá năng lực (ĐGNL) hoặc Đánh giá tư duy (ĐGTD)",
        "mo_ta": (
            "Sử dụng điểm số từ các kỳ thi riêng do các đại học lớn tổ chức "
            "(ĐHQG Hà Nội, ĐHQG TP.HCM, Bách khoa Hà Nội, Sư phạm Hà Nội,…) để xét tuyển."
        ),
    },
    {
        "id": "ket_hop",
        "nhan": "Kết hợp",
        "ten": "Xét tuyển kết hợp / Đề án tuyển sinh riêng của trường",
        "mo_ta": (
            "Kết hợp nhiều tiêu chí khác nhau như chứng chỉ ngoại ngữ quốc tế "
            "(IELTS, TOEFL,…) với kết quả học bạ hoặc điểm thi tốt nghiệp THPT."
        ),
    },
    {
        "id": "xttn_bo",
        "nhan": "Xét thẳng Bộ",
        "ten": "Xét tuyển thẳng và ưu tiên xét tuyển theo quy chế của Bộ GD&ĐT",
        "mo_ta": (
            "Dành cho các đối tượng đặc biệt như anh hùng lao động, chiến sĩ thi đua toàn quốc, "
            "học sinh đoạt giải trong các kỳ thi học sinh giỏi quốc gia, quốc tế, khoa học kỹ thuật, "
            "thể dục thể thao, tay nghề, nghệ thuật cấp quốc gia/quốc tế."
        ),
    },
    {
        "id": "xttn_truong",
        "nhan": "Xét thẳng trường",
        "ten": "Xét tuyển thẳng theo đề án riêng của từng trường",
        "mo_ta": (
            "Áp dụng cho học sinh các trường chuyên, học sinh có chứng chỉ quốc tế uy tín "
            "(SAT, ACT, A-Level) hoặc đạt giải thưởng cấp tỉnh/thành phố theo quy định riêng "
            "của từng trường đại học."
        ),
    },
]

METHOD_IDS = [item["id"] for item in METHOD_CATALOG]

_TOKEN = {
    "hsa": re.compile(r"(?<![a-z0-9])hsa(?![a-z0-9])"),
    "tsa": re.compile(r"(?<![a-z0-9])tsa(?![a-z0-9])"),
    "apt": re.compile(r"(?<![a-z0-9])apt(?![a-z0-9])"),
    "spt": re.compile(r"(?<![a-z0-9])spt(?![a-z0-9])"),
    "sat": re.compile(r"(?<![a-z0-9])sat(?![a-z0-9])"),
    "act": re.compile(r"(?<![a-z0-9])act(?![a-z0-9])"),
    "ib": re.compile(r"(?<![a-z0-9])ib(?![a-z0-9])"),
    "dgnl": re.compile(r"(?<![a-z0-9])dgnl(?![a-z0-9])"),
    "dgtd": re.compile(r"(?<![a-z0-9])dgtd(?![a-z0-9])"),
    "vact": re.compile(r"(?<![a-z0-9])v-?act(?![a-z0-9])"),
}


def _has(text: str, pattern: str) -> bool:
    compiled = _TOKEN.get(pattern)
    if compiled:
        return compiled.search(text) is not None
    return pattern in text


def _split_chunks(text: str) -> List[str]:
    raw = clean_text(text).replace(" • ", ". ").replace(" | ", ". ")
    parts = re.split(r"(?<=[\.\!\?;:])\s+|\n+", raw)
    chunks: List[str] = []
    for part in parts:
        piece = clean_text(part)
        if len(piece) < 12:
            continue
        if len(piece) > 420:
            for sub in re.split(r"\s+(?=\d+[\.\)]\s)|(?<=;)\s+", piece):
                sub = clean_text(sub)
                if len(sub) >= 12:
                    chunks.append(sub[:420])
        else:
            chunks.append(piece)
    return chunks


def _match_ids(folded: str) -> List[str]:
    found: List[str] = []

    exam = any(
        _has(folded, key)
        for key in (
            "ky thi tot nghiep",
            "diem thi thpt",
            "diem thi tot nghiep",
            "ket qua thi tot nghiep",
            "ket qua ky thi tot nghiep",
            "xet diem thi thpt",
            "xet tuyen diem thi",
        )
    )
    if exam or (
        "thpt" in folded
        and any(k in folded for k in ("diem thi", "ket qua thi", "ky thi tot nghiep", "to hop"))
        and "hoc ba" not in folded
    ):
        found.append("thpt")

    if any(
        k in folded
        for k in (
            "xet hoc ba",
            "hoc ba thpt",
            "diem hoc ba",
            "ket qua hoc ba",
            "diem trung binh hoc ba",
        )
    ) or (
        "hoc ba" in folded
        and any(k in folded for k in ("lop 10", "lop 11", "lop 12", "xet tuyen", "ket qua hoc tap"))
    ):
        found.append("hoc_ba")

    if any(
        _has(folded, key)
        for key in ("dgnl", "dgtd", "hsa", "tsa", "apt", "spt", "vact")
    ) or any(k in folded for k in ("danh gia nang luc", "danh gia tu duy")):
        found.append("dgnl")

    language_cert = any(k in folded for k in ("ielts", "toefl", "toeic", "chung chi ngoai ngu", "ccnn"))
    combined = "ket hop" in folded and (
        language_cert
        or "hoc ba" in folded
        or "diem thi" in folded
        or "thpt" in folded
        or "chung chi" in folded
    )
    own_scheme = "de an" in folded and language_cert and (
        "hoc ba" in folded or "thpt" in folded or "ket hop" in folded or "xet tuyen" in folded
    )
    if combined or own_scheme or "xet tuyen ket hop" in folded:
        found.append("ket_hop")

    ministry = any(
        k in folded
        for k in (
            "uu tien xet tuyen",
            "quy che cua bo",
            "quy che bo",
            "bo gd",
            "bo giao duc",
            "anh hung",
            "chien si thi dua",
            "hoc sinh gioi quoc gia",
            "hoc sinh gioi quoc te",
            "giai quoc gia",
            "giai quoc te",
            "olympic",
            "khoa hoc ky thuat quoc gia",
        )
    )
    if ministry or ("xet tuyen thang" in folded and "quy che" in folded and "de an" not in folded):
        found.append("xttn_bo")

    school_direct = any(
        k in folded
        for k in (
            "truong chuyen",
            "hoc sinh chuyen",
            "giai cap tinh",
            "giai tinh",
            "cap tinh",
            "thanh pho",
            "a-level",
            "a level",
        )
    ) or any(_has(folded, key) for key in ("sat", "act", "ib"))
    if ("xet tuyen thang" in folded or "tuyen thang" in folded) and (
        school_direct or ("de an" in folded and "quy che" not in folded)
    ):
        found.append("xttn_truong")
    elif school_direct and any(k in folded for k in ("xet tuyen thang", "tuyen thang", "de an rieng")):
        found.append("xttn_truong")

    return found


def classify_admission_methods(text: str) -> Dict[str, Dict[str, str]]:
    """Trả về {id: {bang_chung}} cho các phương thức tìm thấy trong văn bản."""
    hits: Dict[str, Dict[str, str]] = {}
    for chunk in _split_chunks(text):
        folded = strip_accents(chunk)
        for method_id in _match_ids(folded):
            if method_id in hits:
                continue
            hits[method_id] = {"bang_chung": chunk[:320]}
    return hits


_METHOD_LINE = re.compile(
    r"(?:^|\n)\s*(?:[-–•*]|\d+\.\d+\.?|\d{1,2}[\.\)]|\d{1,2}\s+)\s*"
    r"(Xét tuy[eể]n[^\n]{8,160})",
    re.IGNORECASE,
)
_ADMISSION_CODE = re.compile(
    r"\b((?:EP|POHE|EBBA|EPMP|CLC|TT)\d+|\d{7}(?:_\d+)?|\(\d+\))\b",
    re.IGNORECASE,
)
_MAJOR_CODE = re.compile(r"^\d{7}$")
_CANON_METHODS = (
    ("ket_hop", "Xét tuyển kết hợp", ("ket hop",)),
    ("xttn", "Xét tuyển thẳng", ("tuyen thang", "xet thang")),
    ("thpt", "Xét tuyển theo kết quả thi tốt nghiệp THPT", (
        "tot nghiep thpt", "diem thi thpt", "diem thi tot nghiep", "ptxt5",
    )),
    ("hoc_ba", "Xét tuyển học bạ", ("hoc ba",)),
    ("dgnl", "Xét tuyển bằng kết quả ĐGNL/ĐGTD", (
        "danh gia nang luc", "danh gia tu duy", "dgnl", "dgtd",
    )),
)


def _canon_method(raw: str) -> Optional[tuple]:
    folded = strip_accents(raw)
    if "chinh thuc" in folded or "moi phuong thuc" in folded:
        return None
    for method_id, label, keys in _CANON_METHODS:
        if any(key in folded for key in keys):
            return method_id, label
    return None


def extract_named_methods(text: str) -> List[Dict[str, str]]:
    """Lấy tên hình thức xét tuyển đúng như đề án (xét thẳng, kết hợp, điểm thi THPT, …)."""
    found: Dict[str, Dict[str, str]] = {}
    for match in _METHOD_LINE.finditer(text or ""):
        raw = clean_text(match.group(1))
        canon = _canon_method(raw)
        if not canon or canon[0] in found:
            continue
        summary = re.split(r"\s+(?:ap dung|áp dụng|:)\s+", raw, maxsplit=1, flags=re.I)[0]
        summary = re.sub(r"\s+\d+\s*%$", "", summary).strip(" -–:")
        found[canon[0]] = {
            "id": canon[0],
            "ten": canon[1],
            "mo_ta": summary[:220],
        }
    return list(found.values())


def _program_section(text: str) -> str:
    folded_src = text or ""
    folded = strip_accents(folded_src).lower()
    start = folded.find("ma xet tuyen")
    if start < 0:
        start = folded.find("ten chuong trinh")
    if start < 0:
        start = folded.find("ma nganh")
    if start < 0:
        return ""
    chunk = folded_src[start:]
    end = re.search(r"\bCộng\b|\n\s*5[\.\)]\s", chunk)
    if end:
        chunk = chunk[:end.start()]
    return re.sub(r"\s+", " ", chunk)


def _clean_program_name(value: str) -> str:
    name = clean_text(value)
    name = re.split(r"\s*\(\d+\)\s*", name, maxsplit=1)[0]
    name = re.sub(r"^(?:\d{1,3}\s+)+", "", name)
    name = re.split(
        r"\s+(?:Số|TT|Mã xét|Tên chương trình|Mã ngành|Tên ngành)\b",
        name,
        maxsplit=1,
    )[0]
    while re.search(r"\s+\d{1,4}$", name):
        name = re.sub(r"\s+\d{1,4}$", "", name)
    name = name.strip(" -–|")
    if strip_accents(name).startswith(("ma xet", "ten chuong", "ten nganh")):
        return ""
    return name[:160]


def extract_programs(text: str) -> List[Dict[str, str]]:
    """Bóc mã xét tuyển, tên chương trình, mã ngành và tên ngành trong mục chỉ tiêu."""
    flat = _program_section(text)
    if not flat:
        return []
    marks = list(_ADMISSION_CODE.finditer(flat))
    rows: List[Dict[str, str]] = []
    seen = set()

    def _add(ma_xt: str, prog: str, ma_nganh: str, ten: str, quota: str) -> None:
        prog = _clean_program_name(prog)
        ten = _clean_program_name(ten)
        if len(prog) < 2 or len(ten) < 2 or not _MAJOR_CODE.match(ma_nganh):
            return
        if re.search(r"\d{6,}", prog) or re.search(r"\d{6,}", ten):
            return
        if quota and (not quota.isdigit() or int(quota) > 2000 or int(quota) < 5):
            quota = ""
        key = (ma_xt.upper(), ma_nganh, strip_accents(ten).lower())
        if key in seen:
            return
        seen.add(key)
        rows.append({
            "ma_xet_tuyen": ma_xt.upper(),
            "ten_chuong_trinh": prog,
            "ma_nganh": ma_nganh,
            "ten_nganh": ten,
            "chi_tieu": quota,
        })

    i = 0
    while i < len(marks) - 1:
        token = marks[i].group(1)
        nxt = marks[i + 1].group(1)
        if token.startswith("("):
            i += 1
            continue
        if not _MAJOR_CODE.match(nxt):
            i += 1
            continue
        body = flat[marks[i].end():marks[i + 1].start()]
        after_end = marks[i + 2].start() if i + 2 < len(marks) else len(flat)
        after = flat[marks[i + 1].end():after_end]
        quota_match = re.search(r"\s(\d{2,4})(?:\s|$)", after)
        quota = quota_match.group(1) if quota_match else ""
        ten = after[:quota_match.start()] if quota_match else after
        _add(token, body, nxt, ten, quota)
        i += 2
    return rows


def html_prose(html: str) -> str:
    if not html:
        return ""
    try:
        soup = BeautifulSoup(html, "html.parser")
    except Exception:
        return ""
    for tag in soup(["script", "style", "nav", "footer", "noscript", "svg", "form"]):
        tag.decompose()
    root = None
    for sel in ("article", "main", ".entry-content", ".post-content", "#content", ".description"):
        node = soup.select_one(sel)
        if node and len(node.get_text(strip=True)) > 180:
            root = node
            break
    root = root or soup.body or soup
    for table in root.find_all("table"):
        lines = []
        for tr in table.find_all("tr"):
            cells = [clean_text(cell.get_text(" ", strip=True)) for cell in tr.find_all(["td", "th"])]
            cells = [cell for cell in cells if cell]
            if cells:
                lines.append(" | ".join(cells))
        table.replace_with("\n" + "\n".join(lines) + "\n")
    return root.get_text("\n", strip=True)


def _ocr_pdf(path: str, max_pages: int = 12) -> str:
    """Đọc đề án scan (không có lớp chữ) bằng Vision trên macOS."""
    try:
        import pypdfium2 as pdfium
    except Exception:
        return ""
    swift = os.path.join(os.path.dirname(__file__), "ocr_vision.swift")
    if not os.path.isfile(swift):
        return ""
    import subprocess

    pages: List[str] = []
    pdf = pdfium.PdfDocument(path)
    try:
        page_count = min(len(pdf), max_pages)
    except Exception:
        page_count = 0
    for index in range(page_count):
        image_path = ""
        try:
            bitmap = pdf[index].render(scale=1.4)
            with tempfile.NamedTemporaryFile(suffix=".png", delete=False) as tmp:
                bitmap.to_pil().save(tmp.name)
                image_path = tmp.name
            proc = subprocess.run(
                ["swift", swift, image_path],
                capture_output=True,
                text=True,
                timeout=60,
            )
        except Exception:
            continue
        finally:
            if image_path and os.path.exists(image_path):
                os.unlink(image_path)
        rows: Dict[float, List[tuple]] = {}
        for line in (proc.stdout or "").splitlines():
            parts = line.split("\t", 2)
            if len(parts) != 3:
                continue
            try:
                y = round(float(parts[0]) * 50) / 50
                x = float(parts[1])
            except ValueError:
                continue
            rows.setdefault(y, []).append((x, parts[2]))
        page_lines = [
            " ".join(text for _x, text in sorted(rows[key]))
            for key in sorted(rows, reverse=True)
        ]
        if page_lines:
            pages.append("\n".join(page_lines))
    close = getattr(pdf, "close", None)
    if close:
        close()
    return "\n".join(pages)


def _pdf_prose(path: str, max_pages: int = 12) -> str:
    import pdfplumber

    chunks: List[str] = []
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages[:max_pages]:
            chunks.append(page.extract_text() or "")
    text = "\n".join(chunks)
    if len(clean_text(text)) >= 500:
        return text
    return _ocr_pdf(path, max_pages) or text


def _docx_prose(path: str) -> str:
    import docx

    document = docx.Document(path)
    parts = [p.text for p in document.paragraphs if p.text and p.text.strip()]
    for table in document.tables[:4]:
        for row in table.rows[:40]:
            cells = [clean_text(cell.text) for cell in row.cells if clean_text(cell.text)]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _method_rank(score: int, url: str) -> int:
    blob = _blob(url)
    rank = score
    if "phuong thuc" in blob:
        rank += 80
    if "thong tin tuyen sinh" in blob or "to gap" in blob or "hinh thuc dao tao" in blob:
        rank += 50
    if "de an" in blob or "de-an" in blob or "phuong an tuyen sinh" in blob:
        rank += 40
    if "lien thong" in blob:
        rank -= 40
    if "diem chuan" in blob or "diem-chuan" in blob or "diem trung tuyen" in blob:
        rank -= 70
    return rank


def _empty_methods() -> Dict[str, Dict[str, object]]:
    return {method_id: {"co": False, "bang_chung": "", "nguon": ""} for method_id in METHOD_IDS}


def _merge_hits(
    bucket: Dict[str, Dict[str, object]],
    hits: Dict[str, Dict[str, str]],
    source_url: str,
) -> None:
    for method_id, info in hits.items():
        slot = bucket.setdefault(method_id, {"co": False, "bang_chung": "", "nguon": ""})
        if slot.get("co") and slot.get("bang_chung"):
            continue
        slot["co"] = True
        slot["bang_chung"] = info.get("bang_chung") or ""
        slot["nguon"] = source_url


class AdmissionMethodCollector:
    """Đọc trang phương thức / đề án, không bóc bảng điểm chuẩn."""

    MAX_HTML_PAGES = 10
    MAX_DOCS = 3

    def collect(
        self,
        school_code: str,
        school_name: str,
        website: str,
        years: List[int],
        delay: float = 0.2,
        seed_urls: Optional[List[str]] = None,
        trusted_urls: Optional[List[str]] = None,
    ) -> Dict[str, object]:
        crawler = OfficialSiteCrawler()
        years = sorted({int(y) for y in (years or [])}) or [2026]
        site = normalize_site_url(website)
        if not site:
            for raw in trusted_urls or []:
                site = normalize_site_url(str(raw or "").strip())
                if site:
                    break
        if not site:
            return {
                "ok": False,
                "note": "Không có website chính thức trong danh bạ",
                "records": [],
            }

        by_year: Dict[int, Dict[str, Dict[str, object]]] = {}
        blobs: Dict[int, List[tuple]] = {}
        notes: List[str] = []
        pages_read = 0
        docs_read = 0

        def _bucket(year: int) -> Dict[str, Dict[str, object]]:
            if year not in by_year:
                by_year[year] = _empty_methods()
            return by_year[year]

        def _apply(text: str, source_url: str, year: Optional[int], unspecified: bool) -> None:
            if not clean_text(text):
                return
            target = year if year in years else max(years)
            blobs.setdefault(target, []).append((text, source_url))
            hits = classify_admission_methods(text)
            if hits:
                _merge_hits(_bucket(target), hits, source_url)
            if (unspecified or year not in years) and not notes:
                notes.append(
                    f"Một trang không ghi năm trong các năm đã chọn; gắn với năm {target}."
                )

        base_host = urlparse(site).netloc
        accepted, _rejected = filter_official_urls(site, list(seed_urls or []) + list(trusted_urls or []))
        queue: List = []
        doc_queue: List = []
        seen_html = set()
        try:
            home = crawler._fetch(site, timeout=8)
        except Exception:
            home = None
        if home is not None:
            site = home.url or site
            base_host = urlparse(site).netloc
            home_canon = _canon(site)
            if home_canon:
                seen_html.add(home_canon)
            crawler._harvest(home, site, base_host, years, queue, doc_queue, seen_html)
            home_text = html_prose(home.text or "")
            if home_text:
                _apply(home_text, home.url or site, crawler._page_year(home, years), False)
        for seed in accepted:
            if re.search(r"\.(pdf|docx)(?:$|[?#])", seed, re.I):
                doc_queue.append((80, seed, seed))
            else:
                queue.append((80, seed))
        for portal in admission_portal_urls(site):
            queue.append((60, portal))

        pending = [( _method_rank(score, url), url) for score, url in queue]
        while pending and pages_read < self.MAX_HTML_PAGES:
            pending.sort(key=lambda item: item[0], reverse=True)
            rank, url = pending.pop(0)
            canon = _canon(url)
            if not canon or canon in seen_html or rank < 8:
                continue
            seen_html.add(canon)
            try:
                page = crawler._fetch(url, timeout=12)
            except Exception:
                continue
            if page is None:
                continue
            final_url = page.url or url
            pages_read += 1
            if delay:
                import time
                time.sleep(delay)
            content_type = (page.headers.get("Content-Type") or "").lower()
            if "pdf" in content_type or re.search(r"\.(pdf|docx)(?:$|[?#])", final_url, re.I):
                doc_queue.append((rank, final_url, final_url))
                continue
            prose = html_prose(page.text or "")
            page_year = crawler._page_year(page, years)
            mentioned = [y for y in _years_in(final_url) if y in years]
            unspecified = page_year not in years and not mentioned
            _apply(prose, final_url, page_year if page_year in years else (mentioned[0] if mentioned else None), unspecified)
            crawler._harvest(page, final_url, base_host, years, pending_raw := [], doc_queue, seen_html)
            for score, found in pending_raw:
                pending.append((_method_rank(score, found), found))

        seen_docs = set()
        doc_queue.sort(key=lambda item: _method_rank(item[0], item[1]), reverse=True)
        for score, doc_url, _text in doc_queue:
            if docs_read >= self.MAX_DOCS:
                break
            canon = _canon(doc_url)
            if not canon or canon in seen_docs:
                continue
            if _method_rank(score, doc_url) < 8 and "de an" not in _blob(doc_url) and "phuong thuc" not in _blob(doc_url):
                continue
            seen_docs.add(canon)
            prose, final_url = _document_prose(crawler, doc_url)
            if not prose:
                continue
            docs_read += 1
            doc_years = [y for y in _years_in(final_url, prose[:400]) if y in years]
            year = doc_years[0] if doc_years else None
            _apply(prose, final_url, year, year is None)

        records = []
        years_with_text = sorted(set(by_year) | set(blobs))
        for year in years_with_text:
            combined = "\n".join(text for text, _url in blobs.get(year) or [])
            source_url = ""
            for text, url in blobs.get(year) or []:
                if extract_programs(text) or extract_named_methods(text):
                    source_url = url
                    combined = text if extract_programs(text) else combined
                    if extract_programs(text):
                        break
            programs = extract_programs(combined)
            named = extract_named_methods(combined)
            if not named:
                methods = by_year.get(year) or {}
                named = [
                    {
                        "id": method_id,
                        "ten": next(item["nhan"] for item in METHOD_CATALOG if item["id"] == method_id),
                        "mo_ta": (methods.get(method_id) or {}).get("bang_chung") or "",
                    }
                    for method_id in METHOD_IDS
                    if (methods.get(method_id) or {}).get("co")
                ]
            if not programs and not named:
                continue
            if not source_url and blobs.get(year):
                source_url = blobs[year][0][1]
            note = notes[0] if notes else ""
            if programs and named:
                note = (note + " " if note else "") + "Hình thức trong tài liệu được gắn cho từng ngành."
            base = {
                "ma_truong": school_code,
                "ten_truong": school_name,
                "nam": year,
                "hinh_thuc": named,
                "nguon": source_url,
                "ghi_chu": note.strip(),
            }
            if programs:
                for program in programs:
                    row = dict(base)
                    row.update(program)
                    records.append(row)
            else:
                records.append(base)
        found = []
        for row in records:
            for method in row.get("hinh_thuc") or []:
                label = method.get("ten") or ""
                if label and label not in found:
                    found.append(label)
        host = urlparse(site).netloc or site
        program_count = sum(1 for row in records if row.get("ma_nganh") or row.get("ten_nganh"))
        if records:
            note = f"Website trường {host} — {program_count} ngành, {pages_read} trang, {docs_read} tệp"
        else:
            note = f"Website trường {host} — chưa nhận diện được ngành và hình thức tuyển sinh"
        return {
            "ok": bool(records),
            "note": note,
            "records": records,
            "found": found,
            "programs": program_count,
        }


def _named_from_text(text: str) -> List[Dict[str, str]]:
    named = extract_named_methods(text)
    if named:
        return named
    hits = classify_admission_methods(text)
    return [
        {
            "id": method_id,
            "ten": next(item["nhan"] for item in METHOD_CATALOG if item["id"] == method_id),
            "mo_ta": (hits.get(method_id) or {}).get("bang_chung") or "",
        }
        for method_id in METHOD_IDS
        if method_id in hits
    ]


def _methods_from_cell(value: str) -> List[Dict[str, str]]:
    parts = re.split(r"[;\n|/]+", clean_text(value))
    found: Dict[str, Dict[str, str]] = {}
    for part in parts:
        canon = _canon_method(part)
        if not canon or canon[0] in found:
            continue
        found[canon[0]] = {"id": canon[0], "ten": canon[1], "mo_ta": clean_text(part)[:220]}
    return list(found.values())


def _spreadsheet_rows(path: str) -> List[Dict[str, str]]:
    """Đọc bảng Excel/CSV nếu có cột mã ngành, chương trình hoặc phương thức."""
    import pandas as pd

    ext = os.path.splitext(path)[1].lower()
    frames = []
    if ext == ".csv":
        for encoding in ("utf-8-sig", "utf-8", "cp1258", "latin-1"):
            try:
                frames.append(pd.read_csv(path, header=None, dtype=str, encoding=encoding))
                break
            except Exception:
                continue
    else:
        book = pd.ExcelFile(path)
        for sheet in book.sheet_names:
            try:
                frames.append(pd.read_excel(book, sheet_name=sheet, header=None, dtype=str))
            except Exception:
                continue
    header_keys = {
        "ma_xet_tuyen": ("ma xet tuyen", "ma xt"),
        "ten_chuong_trinh": ("chuong trinh dao tao", "ten chuong trinh"),
        "ma_nganh": ("ma nganh",),
        "ten_nganh": ("ten nganh", "nganh dao tao"),
        "chi_tieu": ("chi tieu",),
        "phuong_thuc": ("phuong thuc", "hinh thuc xet"),
        "ghi_chu": ("ghi chu",),
    }
    parsed: List[Dict[str, str]] = []
    for frame in frames:
        if frame is None or frame.empty:
            continue
        header_at = None
        mapping: Dict[str, int] = {}
        method_cols: List[tuple] = []
        limit = min(len(frame.index), 25)
        for index in range(limit):
            raw_headers = [clean_text(cell) for cell in frame.iloc[index].tolist()]
            labels = [strip_accents(cell).lower() for cell in raw_headers]
            trial: Dict[str, int] = {}
            methods: List[tuple] = []
            for col, label in enumerate(labels):
                if not label:
                    continue
                for field, keys in header_keys.items():
                    if field not in trial and any(key in label for key in keys):
                        trial[field] = col
                        break
                else:
                    if label in {"stt", "tt", "so tt"}:
                        continue
                    methods.append((col, raw_headers[col] if col < len(raw_headers) else label))
            if "ma_nganh" in trial or "ma_xet_tuyen" in trial or "ten_nganh" in trial:
                header_at = index
                mapping = trial
                method_cols = methods
                break
        if header_at is None:
            continue
        for _, series in frame.iloc[header_at + 1:].iterrows():
            cells = [clean_text(cell) for cell in series.tolist()]

            def cell(field: str) -> str:
                col = mapping.get(field)
                text = cells[col] if col is not None and col < len(cells) else ""
                return "" if text.lower() == "nan" else text

            if not any(cell(field) for field in ("ma_xet_tuyen", "ma_nganh", "ten_nganh", "ten_chuong_trinh")):
                continue
            methods = _methods_from_cell(cell("phuong_thuc"))
            known = {item["id"] for item in methods}
            for col, label in method_cols:
                raw = cells[col] if col < len(cells) else ""
                mark = strip_accents(raw).lower()
                if not label or label in known:
                    continue
                known.add(label)
                denied = mark in {"", "khong", "0", "false", "nan", "-", "x", "ko"}
                affirmed = mark in {"co", "1", "true", "v", "yes", "ap dung"}
                combo_list = [] if affirmed or denied else [
                    clean_text(part) for part in re.split(r"[,;]", raw) if clean_text(part)
                ]
                methods.append({
                    "id": label,
                    "ten": label,
                    "ap_dung": not denied,
                    "mo_ta": f"Tổ hợp: {', '.join(combo_list)}" if combo_list else "",
                    "chi_tiet": {"to_hop_xet_tuyen": combo_list} if combo_list else {},
                })
            parsed.append({
                "ma_xet_tuyen": cell("ma_xet_tuyen"),
                "ten_chuong_trinh": cell("ten_chuong_trinh") or cell("ten_nganh"),
                "ma_nganh": cell("ma_nganh"),
                "ten_nganh": cell("ten_nganh") or cell("ten_chuong_trinh"),
                "chi_tieu": re.sub(r"\D", "", cell("chi_tieu"))[:5],
                "ghi_chu": cell("ghi_chu"),
                "hinh_thuc": methods,
            })
    return parsed


def _json_program_rows(path: str) -> List[Dict[str, object]]:
    """Đọc danh sách chương trình: mã, chỉ tiêu, phương thức và ghi chú."""
    with open(path, "r", encoding="utf-8-sig") as handle:
        payload = json.load(handle)
    if isinstance(payload, dict):
        payload = next(
            (
                payload.get(key)
                for key in ("danh_sach", "records", "data", "programs", "chuong_trinh")
                if isinstance(payload.get(key), list)
            ),
            [payload],
        )
    if not isinstance(payload, list):
        raise ValueError("File JSON phải là danh sách chương trình đào tạo")
    rows: List[Dict[str, object]] = []
    for item in payload:
        if not isinstance(item, dict):
            continue
        program = clean_text(str(item.get("chuong_trinh_dao_tao") or item.get("ten_chuong_trinh") or ""))
        code = clean_text(str(item.get("ma") or item.get("ma_xet_tuyen") or ""))
        major_code = clean_text(str(item.get("ma_nganh") or ""))
        major_name = clean_text(str(item.get("ten_nganh") or ""))
        quota = item.get("chi_tieu")
        if quota is None:
            quota = item.get("chi_tieu_du_kien")
        quota_text = ""
        if quota is not None and str(quota).strip() not in {"", "None"}:
            quota_text = re.sub(r"[^\d]", "", str(quota))[:6]
        methods: List[Dict[str, object]] = []
        for method in item.get("phuong_thuc") or []:
            if not isinstance(method, dict):
                continue
            applied = method.get("ap_dung", True)
            enabled = not (
                applied is False or str(applied).strip().lower() in {"false", "0", "khong", "không"}
            )
            name = clean_text(str(method.get("ten") or ""))
            method_code = clean_text(str(method.get("ma") or ""))
            if not name and not method_code:
                continue
            detail = dict(method.get("chi_tiet") or {}) if isinstance(method.get("chi_tiet"), dict) else {}
            combos = method.get("to_hop") or detail.get("to_hop_xet_tuyen") or detail.get("to_hop") or []
            combo_list = []
            if isinstance(combos, list):
                combo_list = [clean_text(str(combo)) for combo in combos if clean_text(str(combo))]
            if combo_list:
                detail["to_hop_xet_tuyen"] = combo_list
            methods.append({
                "id": method_code or name,
                "ten": name or method_code,
                "ap_dung": enabled,
                "mo_ta": f"Tổ hợp: {', '.join(combo_list)}" if combo_list else "",
                "chi_tiet": detail,
            })
        if not any((program, code, major_code, major_name)):
            continue
        rows.append({
            "ma_xet_tuyen": code,
            "ten_chuong_trinh": program,
            "ma_nganh": major_code,
            "ten_nganh": major_name or program,
            "chi_tieu": quota_text,
            "ghi_chu": clean_text(str(item.get("ghi_chu") or "")),
            "hinh_thuc": methods,
        })
    return rows


def local_document_text(path: str) -> str:
    ext = os.path.splitext(path)[1].lower()
    if ext == ".pdf":
        return _pdf_prose(path, max_pages=30)
    if ext == ".docx":
        return _docx_prose(path)
    if ext in {".html", ".htm"}:
        with open(path, "r", encoding="utf-8", errors="replace") as handle:
            return html_prose(handle.read())
    if ext in {".xlsx", ".xls", ".csv"}:
        rows = _spreadsheet_rows(path)
        lines = []
        for row in rows:
            lines.append(" | ".join(
                str(row.get(field) or "")
                for field in ("ma_xet_tuyen", "ten_chuong_trinh", "ma_nganh", "ten_nganh", "chi_tieu")
            ))
        return "\n".join(lines)
    return ""


TMU_ALL_METHODS = (
    ("301", "Xét tuyển thẳng, ưu tiên xét tuyển"),
    ("100", "Thi THPT"),
    ("402", "Xét tuyển theo kết quả HSA, TSA, SAT, ACT"),
    (
        "409",
        "Xét tuyển kết hợp chứng chỉ ngoại ngữ còn hiệu lực tính đến ngày đăng ký xét tuyển với kết quả thi tốt nghiệp THPT",
    ),
    (
        "410",
        "Xét tuyển kết hợp chứng chỉ ngoại ngữ còn hiệu lực tính đến ngày đăng ký xét tuyển với kết quả học tập cấp THPT",
    ),
    (
        "500",
        "Xét tuyển kết hợp giải Nhất, Nhì, Ba trong kỳ thi chọn học sinh giỏi (cấp THPT) cấp tỉnh/thành phố trực",
    ),
)


def _is_tmu_all_method(value: str) -> bool:
    folded = strip_accents(clean_text(value)).lower()
    return "tat ca" in folded and "phuong thuc" in folded


def _expand_tmu_methods(methods: List[Dict[str, object]]) -> List[Dict[str, object]]:
    """TMU ghi «Tất cả các phương thức xét tuyển» nghĩa là đủ 301, 100, 402, 409, 410, 500."""
    kept: List[Dict[str, object]] = []
    shared: List[str] = []
    has_all = False
    for item in methods or []:
        label = str(item.get("ten") or item.get("id") or "")
        if _is_tmu_all_method(label):
            if item.get("ap_dung", True) is not False:
                has_all = True
                detail = item.get("chi_tiet") if isinstance(item.get("chi_tiet"), dict) else {}
                raw = detail.get("to_hop_xet_tuyen") or detail.get("to_hop") or item.get("to_hop") or []
                if isinstance(raw, list):
                    shared = [clean_text(str(code)) for code in raw if clean_text(str(code))]
            continue
        kept.append(item)
    if not has_all:
        return kept
    by_id = {str(item.get("id") or item.get("ten") or ""): item for item in kept}
    combo_note = f"Tổ hợp: {', '.join(shared)}" if shared else ""
    for code, note in TMU_ALL_METHODS:
        current = by_id.get(code)
        if current:
            current["ap_dung"] = True
            detail = current.get("chi_tiet") if isinstance(current.get("chi_tiet"), dict) else {}
            if shared and not (detail.get("to_hop_xet_tuyen") or detail.get("to_hop")):
                current["chi_tiet"] = {**detail, "to_hop_xet_tuyen": shared}
                current["mo_ta"] = combo_note
            continue
        kept.append({
            "id": code,
            "ten": code,
            "ap_dung": True,
            "mo_ta": combo_note or note,
            "chi_tiet": {"to_hop_xet_tuyen": list(shared)} if shared else {},
        })
    return kept


def records_from_local_file(
    path: str,
    school_code: str,
    school_name: str,
    year: int,
    source_url: str = "",
    filename: str = "",
) -> List[Dict[str, object]]:
    """Bóc ngành và hình thức xét tuyển từ tài liệu người dùng tải lên."""
    ext = os.path.splitext(path)[1].lower()
    if ext == ".json":
        sheet_rows = _json_program_rows(path)
    else:
        sheet_rows = _spreadsheet_rows(path) if ext in {".xlsx", ".xls", ".csv"} else []
    prose = "" if sheet_rows else local_document_text(path)
    programs = sheet_rows or extract_programs(prose)
    named = _named_from_text(prose) if prose else []
    if not programs and not named:
        return []
    note = f"Tải lên từ {filename}" if filename else "Tải lên từ tài liệu"
    base = {
        "ma_truong": school_code,
        "ten_truong": school_name,
        "nam": int(year),
        "hinh_thuc": named,
        "nguon": source_url,
        "ghi_chu": note,
    }
    records = []
    if programs:
        for program in programs:
            row = dict(base)
            methods = program.get("hinh_thuc") or named
            row.update({key: value for key, value in program.items() if key != "hinh_thuc"})
            row["hinh_thuc"] = methods
            if program.get("ghi_chu"):
                row["ghi_chu"] = clean_text(str(program.get("ghi_chu")))
            records.append(row)
    else:
        records.append(base)
    if str(school_code or "").upper() == "TMU":
        for row in records:
            row["hinh_thuc"] = _expand_tmu_methods(row.get("hinh_thuc") or [])
    return records


def _document_prose(crawler: OfficialSiteCrawler, url: str):
    res = crawler._fetch(url, timeout=20)
    if res is None or len(res.content) > crawler.MAX_DOC_BYTES:
        return "", url
    path = urlparse(res.url or url).path.lower()
    content_type = (res.headers.get("Content-Type") or "").lower()
    if path.endswith(".pdf") or "pdf" in content_type or res.content[:4] == b"%PDF":
        suffix = ".pdf"
        reader = _pdf_prose
    elif path.endswith(".docx"):
        suffix = ".docx"
        reader = _docx_prose
    else:
        return "", res.url or url
    tmp_path = ""
    try:
        with tempfile.NamedTemporaryFile(suffix=suffix, delete=False) as tmp:
            tmp.write(res.content)
            tmp_path = tmp.name
        return reader(tmp_path), res.url or url
    except Exception:
        return "", res.url or url
    finally:
        if tmp_path and os.path.exists(tmp_path):
            os.unlink(tmp_path)
