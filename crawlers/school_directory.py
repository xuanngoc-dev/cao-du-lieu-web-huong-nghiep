# -*- coding: utf-8 -*-
"""
Module: crawlers.school_directory
Mô tả: Đọc danh bạ mã trường và website chính thức từ cấu hình cục bộ.
"""

import os
import re
import json
from concurrent.futures import ThreadPoolExecutor, as_completed
from typing import Dict, List, Optional, Any, Tuple

import requests
from bs4 import BeautifulSoup

from core.normalizer import clean_text


# Nhãn loại hình chuẩn
TYPE_DAI_HOC = "dai_hoc"
TYPE_CAO_DANG = "cao_dang"
TYPE_HOC_VIEN = "hoc_vien"

TYPE_LABELS = {
    TYPE_DAI_HOC: "Đại học",
    TYPE_CAO_DANG: "Cao đẳng",
    TYPE_HOC_VIEN: "Học viện",
}

# Tỉnh/thành theo miền. Khớp cả tên cũ còn trong địa chỉ.
_REGION_PROVINCES = {
    "bac": (
        "ha noi", "hai phong", "quang ninh", "bac ninh", "bac giang", "ha nam",
        "nam dinh", "thai binh", "ninh binh", "vinh phuc", "phu tho", "hoa binh",
        "son la", "dien bien", "lai chau", "lao cai", "yen bai", "tuyen quang",
        "ha giang", "cao bang", "bac kan", "lang son", "thai nguyen", "hung yen",
        "hai duong",
    ),
    "trung": (
        "thanh hoa", "nghe an", "ha tinh", "quang binh", "quang tri",
        "thua thien hue", "hue", "da nang", "quang nam", "quang ngai", "binh dinh",
        "phu yen", "khanh hoa", "ninh thuan", "binh thuan", "kon tum", "gia lai",
        "dak lak", "dak nong", "dac lac", "dac nong",
    ),
    "nam": (
        "ho chi minh", "sai gon", "tp hcm", "binh duong", "binh phuoc", "dong nai",
        "tay ninh", "ba ria", "vung tau", "long an", "tien giang", "ben tre",
        "tra vinh", "vinh long", "dong thap", "an giang", "kien giang", "can tho",
        "hau giang", "soc trang", "bac lieu", "ca mau", "lam dong",
    ),
}

_SECTOR_RULES = (
    ("Y - Dược", ("y duoc", "y khoa", "y te", " duoc", "dieu duong", " y ")),
    ("Sư phạm", ("su pham", "giao duc")),
    ("Luật", ("luat",)),
    ("Kinh tế", ("kinh te", "thuong mai", "tai chinh", "ngan hang", "ngoai thuong", "ke toan")),
    ("Kỹ thuật", ("bach khoa", "ky thuat", "cong nghe", "cong nghiep", "xay dung", "giao thong", "mo dia chat", "dien luc", "thuy loi")),
    ("Nông - Lâm - Ngư", ("nong nghiep", "lam nghiep", "thuy san")),
    ("Nghệ thuật", ("my thuat", "am nhac", "san khau", "dien anh", "van hoa nghe thuat")),
    ("Ngoại ngữ", ("ngoai ngu",)),
    ("An ninh - Quốc phòng", ("cong an", "an ninh", "quan su", "bien phong", "hau can")),
    ("Du lịch", ("du lich",)),
    ("Báo chí - Truyền thông", ("bao chi", "truyen thong")),
    ("Thể dục - Thể thao", ("the duc", "the thao")),
)


def _fold_match(text: str) -> str:
    from core.normalizer import strip_accents
    folded = strip_accents(text or "")
    folded = folded.replace("tp.", "tp ").replace("t.p", "tp")
    return re.sub(r"[^a-z0-9]+", " ", folded).strip()


def _region_of_text(text: str) -> str:
    folded = f" {_fold_match(text)} "
    if " ho chi minh " in folded or " sai gon " in folded or " tp hcm " in folded:
        return "nam"
    found = ""
    found_at = -1
    for region, names in _REGION_PROVINCES.items():
        for name in names:
            at = folded.rfind(f" {name} ")
            if at > found_at:
                found = region
                found_at = at
    return found


def split_addresses_by_region(dia_chi: str) -> Dict[str, str]:
    """Tách địa chỉ thành cơ sở miền Bắc, Trung, Nam."""
    raw = clean_text(dia_chi or "")
    result = {"dia_chi_bac": "", "dia_chi_trung": "", "dia_chi_nam": ""}
    if not raw:
        return result
    chunks = [
        clean_text(part).strip(" -")
        for part in re.split(
            r"(?:\s+-\s+-\s+|\n+|;\s*(?=cơ sở)|\s+(?=cơ sở\b))",
            raw,
            flags=re.I,
        )
        if clean_text(part).strip(" -")
    ]
    if not chunks:
        chunks = [raw]
    buckets = {"bac": [], "trung": [], "nam": []}
    unknown = []
    for chunk in chunks:
        region = _region_of_text(chunk)
        if region:
            buckets[region].append(chunk)
        else:
            unknown.append(chunk)
    if unknown and sum(bool(v) for v in buckets.values()) == 1:
        only = next(key for key, val in buckets.items() if val)
        buckets[only].extend(unknown)
    elif unknown and not any(buckets.values()):
        return result
    result["dia_chi_bac"] = "\n".join(dict.fromkeys(buckets["bac"]))
    result["dia_chi_trung"] = "\n".join(dict.fromkeys(buckets["trung"]))
    result["dia_chi_nam"] = "\n".join(dict.fromkeys(buckets["nam"]))
    labels = []
    if result["dia_chi_bac"]:
        labels.append("Bắc")
    if result["dia_chi_trung"]:
        labels.append("Trung")
    if result["dia_chi_nam"]:
        labels.append("Nam")
    result["khu_vuc"] = ", ".join(labels)
    return result


def classify_school_sector(name: str) -> str:
    """Loại trường theo lĩnh vực: kinh tế, kỹ thuật, y dược, …"""
    folded = f" {_fold_match(name)} "
    for label, hints in _SECTOR_RULES:
        if any(hint in folded for hint in hints):
            return label
    return "Đa ngành"

# Khóa hồ sơ trường (từ mục Giới thiệu trên trang đề án)
PROFILE_KEYS = [
    "thong_tin_chung",
    "dia_chi",
    "website",
    "hotline",
    "fanpage",
    "linh_vuc_chuong_trinh",
    "vi_the_thanh_tuu",
    "gioi_thieu_url",
]

# Liên hệ cơ bản — luôn lấy kể cả khi tắt hồ sơ giới thiệu đầy đủ
CONTACT_KEYS = ["dia_chi", "website", "hotline", "fanpage", "gioi_thieu_url"]

# Nguồn thông báo tuyển sinh trên website chính thức của trường
NOTICE_KEYS = ["domain_diem_chuan", "link_quy_che"]


def _origin(url: str) -> str:
    from urllib.parse import urlparse

    parsed = urlparse(url or "")
    if parsed.scheme not in ("http", "https") or not parsed.netloc:
        return ""
    return f"{parsed.scheme}://{parsed.netloc.lower()}"


def discover_admission_notices(website: str, timeout: int = 8) -> Dict[str, str]:
    """
    Tìm domain đăng thông báo điểm chuẩn và link thông báo quy chế tuyển sinh
    trên website chính thức của trường (kể cả cổng tuyển sinh cùng tổ chức).
    """
    from urllib.parse import urljoin, urlparse

    from crawlers.official_site_crawler import (
        _blob,
        _same_org,
        admission_portal_urls,
        normalize_site_url,
    )

    empty = {key: "" for key in NOTICE_KEYS}
    site = normalize_site_url(website)
    if not site:
        return empty

    session = requests.Session()
    session.headers.update(SchoolDirectory.HEADERS)
    org_host = urlparse(site).netloc

    def fetch(url: str, wait: int) -> Optional[requests.Response]:
        try:
            res = session.get(url, timeout=wait, allow_redirects=True)
        except requests.exceptions.SSLError:
            try:
                res = session.get(url, timeout=wait, allow_redirects=True, verify=False)
            except requests.exceptions.RequestException:
                return None
        except requests.exceptions.RequestException:
            return None
        if res.status_code != 200:
            return None
        if not _same_org(res.url or url, org_host):
            return None
        kind = (res.headers.get("Content-Type") or "").lower()
        if "html" not in kind and "text" not in kind:
            return None
        return res

    pages: List[requests.Response] = []
    home = fetch(site, timeout)
    if home is not None:
        pages.append(home)
    for portal in admission_portal_urls(site):
        if sum(1 for page in pages if page is not home) >= 2:
            break
        found = fetch(portal, min(timeout, 6))
        if found is None:
            continue
        sample = _blob(found.url or "", (found.text or "")[:4000])
        if any(hint in sample for hint in ("tuyen sinh", "diem chuan", "xet tuyen", "quy che", "de an")):
            pages.append(found)

    cutoff_best = ("", 0)
    regulation_best = ("", 0)
    hub_best = ("", 0)

    def consider(url: str, label: str) -> None:
        nonlocal cutoff_best, regulation_best, hub_best
        folded = _blob(label, url)
        if not folded:
            return
        penalty = 0
        for hint, weight in (
            ("du bao", 18),
            ("thac si", 20),
            ("tien si", 16),
            ("tuyen dung", 16),
            ("lien thong", 12),
            ("lien ket", 12),
            ("cao hoc", 12),
        ):
            if hint in folded:
                penalty += weight
        year_bonus = 0
        for year, bonus in (("2026", 8), ("2025", 5), ("2024", 2)):
            if year in folded:
                year_bonus = bonus
                break
        cutoff = year_bonus - penalty
        for hint, weight in (
            ("diem chuan", 22),
            ("diem trung tuyen", 18),
            ("thong bao diem", 14),
            ("cong bo diem", 12),
        ):
            if hint in folded:
                cutoff += weight
        if "quy doi" in folded:
            cutoff -= 8
        if cutoff > cutoff_best[1]:
            cutoff_best = (url, cutoff)

        regulation = year_bonus - penalty
        for hint, weight in (
            ("quy che tuyen sinh", 32),
            ("de an tuyen sinh", 28),
            ("phuong an tuyen sinh", 24),
            ("thong bao tuyen sinh", 16),
            ("quy che", 10),
            ("de an", 6),
        ):
            if hint in folded:
                regulation += weight
        if regulation > regulation_best[1]:
            regulation_best = (url, regulation)

        hub = 0
        if any(hint in folded for hint in ("tuyen sinh", "xet tuyen", "thong bao")):
            hub = 8 + year_bonus - penalty
        if hub > hub_best[1]:
            hub_best = (url, hub)

    seen_pages = set()
    for page in pages:
        final = page.url or ""
        if final in seen_pages:
            continue
        seen_pages.add(final)
        consider(final, "")
        try:
            soup = BeautifulSoup(page.text or "", "html.parser")
        except Exception:
            continue
        for anchor in soup.find_all("a", href=True):
            href = (anchor.get("href") or "").strip()
            if not href or href.startswith(("#", "mailto:", "tel:", "javascript:")):
                continue
            absolute = urljoin(final, href)
            if not _same_org(absolute, org_host):
                continue
            consider(absolute, clean_text(anchor.get_text(" ", strip=True)))

    if regulation_best[1] < 18 and hub_best[1] >= 8 and hub_best[0] not in seen_pages:
        extra = fetch(hub_best[0], timeout)
        if extra is not None:
            final = extra.url or hub_best[0]
            consider(final, "")
            try:
                soup = BeautifulSoup(extra.text or "", "html.parser")
                for anchor in soup.find_all("a", href=True):
                    href = (anchor.get("href") or "").strip()
                    if not href or href.startswith(("#", "mailto:", "tel:", "javascript:")):
                        continue
                    absolute = urljoin(final, href)
                    if _same_org(absolute, org_host):
                        consider(absolute, clean_text(anchor.get_text(" ", strip=True)))
            except Exception:
                pass

    if cutoff_best[1] >= 12:
        empty["domain_diem_chuan"] = _origin(cutoff_best[0])
    else:
        portal_pages = [p for p in pages if _origin(p.url or "") != _origin(site)]
        if portal_pages:
            empty["domain_diem_chuan"] = _origin(portal_pages[0].url or "")
    if regulation_best[1] >= 24:
        empty["link_quy_che"] = regulation_best[0]
    return empty


def _has_contact(info: dict) -> bool:
    """Đủ website + địa chỉ để hiển thị bảng danh bạ."""
    return bool((info or {}).get("website")) and bool((info or {}).get("dia_chi"))


def _has_full_profile(info: dict) -> bool:
    return bool(
        (info or {}).get("thong_tin_chung")
        or (info or {}).get("vi_the_thanh_tuu")
        or (info or {}).get("linh_vuc_chuong_trinh")
    )

_TRAINING_HINTS = (
    "đào tạo", "dao tao", "lĩnh vực", "linh vuc", "chương trình", "chuong trinh",
    "ngành", "nganh", "khối ngành", "chuyên ngành",
)
_ACHIEVE_HINTS = (
    "sứ mạng", "su mang", "tầm nhìn", "tam nhin", "thành tựu", "thanh tuu",
    "vị thế", "vi the", "xếp hạng", "hang đầu", "hàng đầu", "thành lập",
    "giá trị cốt lõi", "triết lý", "phương châm", "lịch sử", "lich su",
)


def _split_label_value(text: str) -> Tuple[str, str]:
    """Tách 'Địa chỉ : Số 1…' → ('địa chỉ', 'Số 1…')."""
    t = clean_text(text or "")
    if not t:
        return "", ""
    for sep in (":", "："):
        if sep in t:
            left, right = t.split(sep, 1)
            return clean_text(left).lower(), clean_text(right)
    return "", t


def _is_training_para(text: str) -> bool:
    low = (text or "").lower()
    return any(k in low for k in _TRAINING_HINTS)


def _is_achieve_para(text: str) -> bool:
    low = (text or "").lower()
    return any(k in low for k in _ACHIEVE_HINTS)


def classify_school_type(name: str, page_hint: str = "") -> str:
    """Phân loại loại hình trường theo tên và nguồn trang."""
    if page_hint == TYPE_CAO_DANG:
        return TYPE_CAO_DANG
    n = (name or "").lower()
    if "cao đẳng" in n or "cao dang" in n:
        return TYPE_CAO_DANG
    if "học viện" in n or "hoc vien" in n:
        return TYPE_HOC_VIEN
    return TYPE_DAI_HOC


def extract_code_from_slug_and_text(slug: str, text: str) -> str:
    """
    Trích mã trường từ slug hoặc text liên kết.
    Hỗ trợ mã ĐH (BKA, KHA) và mã CĐ dài (CDD0209, CDT026, 063).
    """
    code = ""
    # Ưu tiên mã ở đầu text: "BKA-Đại học...", "CDD0209-Trường Cao đẳng..."
    match_text = re.match(r"^([A-Za-z0-9]{2,10})[\s\-:]", text.strip())
    if match_text:
        code = match_text.group(1).upper()
    if not code:
        # Mã ở cuối slug: ...-BKA, ...-CDD0209, ...-063
        match_slug = re.search(r"-([A-Za-z0-9]{2,10})$", slug)
        if match_slug:
            code = match_slug.group(1).upper()
    return code


class SchoolDirectory:
    """
    Danh bạ mã trường Đại học & Cao đẳng.
    Mọi URL dùng để thu thập dữ liệu đều là website chính thức của trường.
    """

    BASE_URL = ""
    HEADERS = {
        "User-Agent": (
            "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) "
            "AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36"
        )
    }

    LISTING_SOURCES = [
        {"url": "/diem-chuan.html", "hint": TYPE_DAI_HOC, "label": "Đại học / Học viện"},
        {"url": "/diem-chuan.html?type=cao-dang", "hint": TYPE_CAO_DANG, "label": "Cao đẳng"},
    ]

    def __init__(self, cache_dir: str = "data/cache"):
        self.cache_dir = cache_dir
        os.makedirs(self.cache_dir, exist_ok=True)
        self.cache_file = os.path.join(self.cache_dir, "school_slugs_cache.json")
        self.schools: Dict[str, Dict[str, Any]] = {}

    def _read_cache_file(self) -> Dict[str, Dict[str, Any]]:
        if not os.path.exists(self.cache_file):
            return {}
        try:
            with open(self.cache_file, "r", encoding="utf-8") as f:
                data = json.load(f)
            return data if isinstance(data, dict) else {}
        except Exception:
            return {}

    def _read_official_catalog(self) -> Dict[str, Dict[str, Any]]:
        path = os.path.join(
            os.path.dirname(os.path.dirname(__file__)),
            "config",
            "schools_all.json",
        )
        try:
            with open(path, "r", encoding="utf-8") as f:
                raw = json.load(f)
        except (OSError, json.JSONDecodeError, TypeError):
            return {}
        schools: Dict[str, Dict[str, Any]] = {}
        for item in raw.get("schools") or []:
            if not isinstance(item, dict):
                continue
            code = str(item.get("code") or "").strip().upper()
            if not code:
                continue
            entry = dict(item)
            entry["code"] = code
            entry["slug"] = entry.get("online_slug") or entry.get("slug") or ""
            entry["type"] = entry.get("type") or classify_school_type(
                entry.get("name") or ""
            )
            entry["type_label"] = TYPE_LABELS.get(entry["type"], entry["type"])
            # Không phát sinh link giới thiệu từ nguồn tổng hợp bên ngoài.
            entry["gioi_thieu_url"] = entry.get("website") or ""
            schools[code] = entry
        return schools

    @staticmethod
    def _merge_preserved_fields(
        fetched: Dict[str, Dict[str, Any]],
        previous: Dict[str, Dict[str, Any]],
    ) -> Dict[str, Dict[str, Any]]:
        """Giữ website/địa chỉ/hồ sơ cũ khi làm mới danh bạ từ web."""
        if not previous:
            return fetched
        for code, entry in fetched.items():
            prev = previous.get(code)
            if not isinstance(prev, dict):
                continue
            for key in PROFILE_KEYS:
                if not entry.get(key) and prev.get(key):
                    entry[key] = prev[key]
        return fetched

    def load(
        self,
        force_refresh: bool = False,
        include_dai_hoc: bool = True,
        include_cao_dang: bool = True,
        include_profile: bool = False,
        include_contact: bool = True,
        profile_limit: int = 0,
    ) -> Dict[str, Dict[str, Any]]:
        """Nạp danh bạ cục bộ; không gọi cổng dữ liệu bên ngoài."""
        catalog = self._read_official_catalog()
        self.schools = self._filter_types(
            catalog, include_dai_hoc, include_cao_dang
        )
        return self.schools

        # Mã cũ bên dưới chỉ giữ để đọc các file cache lịch sử; không còn được gọi.
        previous_cache = self._read_cache_file()
        if not force_refresh and previous_cache:
            try:
                data = previous_cache
                # Cache cũ (không có type) → gắn type suy luận
                normalized = {}
                for code, info in data.items():
                    if not isinstance(info, dict):
                        continue
                    item = dict(info)
                    item.setdefault("code", code)
                    item["type"] = item.get("type") or classify_school_type(
                        item.get("name", ""), ""
                    )
                    item["type_label"] = TYPE_LABELS.get(item["type"], item["type"])
                    normalized[code] = item
                # Nếu cache chỉ có ĐH (không có CĐ) và user muốn CĐ → refresh
                has_cd = any(v.get("type") == TYPE_CAO_DANG for v in normalized.values())
                if include_cao_dang and not has_cd:
                    print("[DANH_BA] Cache chưa có Cao đẳng → làm mới danh bạ...")
                else:
                    self.schools = self._filter_types(
                        normalized, include_dai_hoc, include_cao_dang
                    )
                    full = dict(normalized)
                    if include_profile:
                        self.enrich_profiles(full, only_missing=True, limit=profile_limit)
                        self._save_cache(full)
                    elif include_contact:
                        self.enrich_profiles(
                            full, only_missing=True, contact_only=True, limit=profile_limit
                        )
                        self._save_cache(full)
                    self.schools = self._filter_types(
                        full, include_dai_hoc, include_cao_dang
                    )
                    return self.schools
            except Exception:
                pass

        fetched = self.fetch_from_web(
            include_dai_hoc=include_dai_hoc,
            include_cao_dang=include_cao_dang,
        )
        fetched = self._merge_preserved_fields(fetched, previous_cache)
        if include_profile:
            self.enrich_profiles(fetched, only_missing=False, limit=profile_limit)
        elif include_contact:
            self.enrich_profiles(
                fetched, only_missing=True, contact_only=True, limit=profile_limit
            )
        self.schools = fetched
        self._save_cache(fetched)
        return self.schools

    def fetch_from_web(
        self,
        include_dai_hoc: bool = True,
        include_cao_dang: bool = True,
    ) -> Dict[str, Dict[str, Any]]:
        """Tương thích API cũ: trả danh bạ cục bộ, không truy cập nguồn bên ngoài."""
        return self._filter_types(
            self._read_official_catalog(), include_dai_hoc, include_cao_dang
        )

        # Mã parser lịch sử bên dưới không còn được gọi.
        directory: Dict[str, Dict[str, Any]] = {}

        for src in self.LISTING_SOURCES:
            hint = src["hint"]
            if hint == TYPE_DAI_HOC and not include_dai_hoc:
                continue
            if hint == TYPE_CAO_DANG and not include_cao_dang:
                continue

            url = self.BASE_URL + src["url"]
            print(f"[DANH_BA] Đang lấy danh sách {src['label']}: {url}")
            try:
                res = requests.get(url, headers=self.HEADERS, timeout=20)
                if res.status_code != 200:
                    print(f"[CẢNH BÁO] HTTP {res.status_code} khi tải {url}")
                    continue
                soup = BeautifulSoup(res.text, "html.parser")
                count_before = len(directory)

                for a in soup.find_all("a", href=True):
                    href = a["href"]
                    if "/diem-chuan/" not in href or not href.endswith(".html"):
                        continue
                    text = clean_text(a.get_text())
                    slug = href.replace("/diem-chuan/", "").replace(".html", "")
                    if not slug or slug in ("index",):
                        continue

                    code = extract_code_from_slug_and_text(slug, text)
                    if not code:
                        continue

                    clean_name = re.sub(
                        r"^[A-Za-z0-9]{2,10}[\s\-:]+", "", text
                    ).strip() or text
                    school_type = classify_school_type(clean_name, page_hint=hint)

                    # Không ghi đè ĐH bằng bản ghi CĐ trùng mã (hiếm), ưu tiên bản ghi đầy đủ hơn
                    existing = directory.get(code)
                    entry = {
                        "code": code,
                        "name": clean_name,
                        "slug": slug,
                        "type": school_type,
                        "type_label": TYPE_LABELS.get(school_type, school_type),
                        "listing_url": url,
                    }
                    if existing:
                        # Giữ tên dài hơn / slug nếu đã có
                        if len(existing.get("name", "")) > len(clean_name):
                            entry["name"] = existing["name"]
                        if school_type == TYPE_CAO_DANG or existing.get("type") == TYPE_DAI_HOC:
                            # Ưu tiên type từ trang đúng hơn
                            if hint == TYPE_CAO_DANG:
                                entry["type"] = TYPE_CAO_DANG
                                entry["type_label"] = TYPE_LABELS[TYPE_CAO_DANG]
                    directory[code] = entry

                added = len(directory) - count_before
                print(f"[DANH_BA] +{added} trường từ {src['label']} (tổng: {len(directory)})")
            except Exception as e:
                print(f"[CẢNH BÁO] Lỗi khi lấy danh bạ {src['label']}: {e}")

        print(f"[DANH_BA] Hoàn tất: {len(directory)} trường (ĐH/HV/CĐ).")
        return directory

    def fetch_school_profile(
        self,
        slug: str,
        code: str = "",
        contact_only: bool = False,
    ) -> Dict[str, str]:
        """Trả hồ sơ đã lưu cùng website chính thức; không gọi nguồn tổng hợp."""
        empty = {k: "" for k in PROFILE_KEYS}
        catalog = self._read_official_catalog()
        item = catalog.get((code or "").strip().upper())
        if item is None and slug:
            item = next(
                (entry for entry in catalog.values() if entry.get("slug") == slug),
                None,
            )
        if not item:
            return empty
        return {
            key: (
                item.get("website") if key == "gioi_thieu_url"
                else item.get(key)
            ) or ""
            for key in PROFILE_KEYS
        }

        # Parser lịch sử bên dưới không còn được gọi.
        url = f"{self.BASE_URL}/de-an-tuyen-sinh/{slug}.html"
        empty["gioi_thieu_url"] = url + "#gioi-thieu"
        try:
            res = requests.get(url, headers=self.HEADERS, timeout=18)
            if res.status_code != 200:
                return empty
            soup = BeautifulSoup(res.text, "html.parser")
            section = soup.find(id="gioi-thieu")
            if not section:
                # Fallback: trang gioi-thieu-truong-{slug}
                alt = f"{self.BASE_URL}/gioi-thieu-truong-{slug}.html"
                res2 = requests.get(alt, headers=self.HEADERS, timeout=15)
                if res2.status_code == 200:
                    soup = BeautifulSoup(res2.text, "html.parser")
                    section = soup.find(id="gioi-thieu") or soup
                    empty["gioi_thieu_url"] = alt
            if not section:
                return empty

            info_box = section.select_one(".basic-info__info") or section
            achieve_box = None if contact_only else section.select_one(".basic-info__achievement")

            info_lines: List[str] = []
            dia_chi_parts: List[str] = []
            website = hotline = fanpage = ""

            for li in info_box.find_all("li"):
                raw = clean_text(li.get_text(" ", strip=True))
                if not raw:
                    continue
                label, value = _split_label_value(raw)
                info_lines.append(raw)
                if not value:
                    continue
                if any(k in label for k in ("địa chỉ", "dia chi", "cơ sở", "co so", "campus")):
                    dia_chi_parts.append(value if ":" not in raw else value)
                elif "website" in label or "web site" in label:
                    website = value
                elif any(k in label for k in ("hotline", "điện thoại", "dien thoai", "tel", "phone")):
                    hotline = value
                elif any(k in label for k in ("fanpage", "facebook", "fb")):
                    fanpage = value

            # Hotline đôi khi nằm trong cùng dòng địa chỉ / mô tả
            if not hotline:
                blob = " ".join(info_lines)
                m = re.search(
                    r"(?:hotline|điện thoại|tel)\s*[:\-]?\s*([0-9.\s\-–+/]{8,})",
                    blob,
                    re.IGNORECASE,
                )
                if m:
                    hotline = clean_text(m.group(1))

            if contact_only:
                return {
                    "thong_tin_chung": "",
                    "dia_chi": "\n".join(dict.fromkeys(dia_chi_parts)),
                    "website": website,
                    "hotline": hotline,
                    "fanpage": fanpage,
                    "linh_vuc_chuong_trinh": "",
                    "vi_the_thanh_tuu": "",
                    "gioi_thieu_url": empty["gioi_thieu_url"],
                }

            training_parts: List[str] = []
            achieve_parts: List[str] = []
            if achieve_box:
                paras = []
                for tag in achieve_box.find_all(["p", "li"]):
                    t = clean_text(tag.get_text(" ", strip=True))
                    if t and len(t) > 15:
                        paras.append(t)
                if not paras:
                    t = clean_text(achieve_box.get_text("\n", strip=True))
                    paras = [p.strip() for p in t.split("\n") if len(p.strip()) > 15]

                for p in paras:
                    if _is_training_para(p) and not _is_achieve_para(p):
                        training_parts.append(p)
                    elif _is_achieve_para(p):
                        achieve_parts.append(p)
                    elif _is_training_para(p):
                        training_parts.append(p)
                    else:
                        # Đoạn lịch sử / giới thiệu chung → vị thế & thành tựu
                        achieve_parts.append(p)

            # Nếu không tách được, đổ toàn bộ achievement vào vị thế
            if achieve_box and not training_parts and not achieve_parts:
                achieve_parts.append(clean_text(achieve_box.get_text(" ", strip=True))[:4000])

            return {
                "thong_tin_chung": "\n".join(info_lines),
                "dia_chi": "\n".join(dict.fromkeys(dia_chi_parts)),
                "website": website,
                "hotline": hotline,
                "fanpage": fanpage,
                "linh_vuc_chuong_trinh": "\n\n".join(training_parts)[:5000],
                "vi_the_thanh_tuu": "\n\n".join(achieve_parts)[:5000],
                "gioi_thieu_url": empty["gioi_thieu_url"],
            }
        except Exception as e:
            print(f"[CẢNH BÁO] Không lấy được hồ sơ {code or slug}: {e}")
            return empty

    def enrich_profiles(
        self,
        schools: Optional[Dict[str, Dict[str, Any]]] = None,
        only_missing: bool = True,
        max_workers: int = 6,
        limit: int = 0,
        contact_only: bool = False,
    ) -> Dict[str, Dict[str, Any]]:
        """
        Bổ sung hồ sơ / liên hệ cho danh bạ (song song, có cache từng trường).
        contact_only=True: chỉ điền website + địa chỉ (và hotline/fanpage nếu có).
        """
        target = schools if schools is not None else self.schools
        if not target:
            return target

        items = list(target.items())
        if only_missing:
            if contact_only:
                items = [(c, inf) for c, inf in items if not _has_contact(inf)]
            else:
                items = [
                    (c, inf) for c, inf in items
                    if not (_has_full_profile(inf) or _has_contact(inf))
                ]
        if limit and limit > 0:
            items = items[:limit]

        if not items:
            label = "Website/địa chỉ" if contact_only else "Hồ sơ giới thiệu"
            print(f"[DANH_BA] {label} đã có sẵn — bỏ qua bước bổ sung.")
            return target

        label = "website/địa chỉ" if contact_only else "hồ sơ trường"
        print(f"[DANH_BA] Đang lấy {label} cho {len(items)} trường…")
        done = 0

        def _job(code_info):
            code, info = code_info
            profile = self.fetch_school_profile(
                info.get("slug") or "", code=code, contact_only=contact_only
            )
            return code, profile

        with ThreadPoolExecutor(max_workers=max_workers) as pool:
            futures = [pool.submit(_job, it) for it in items]
            for fut in as_completed(futures):
                code, profile = fut.result()
                if code in target:
                    if contact_only:
                        for key in CONTACT_KEYS:
                            val = profile.get(key) or ""
                            if val and not target[code].get(key):
                                target[code][key] = val
                    else:
                        target[code].update(profile)
                if code in self.schools:
                    if contact_only:
                        for key in CONTACT_KEYS:
                            val = profile.get(key) or ""
                            if val and not self.schools[code].get(key):
                                self.schools[code][key] = val
                    else:
                        self.schools[code].update(profile)
                done += 1
                if done % 25 == 0 or done == len(items):
                    print(f"[DANH_BA] {label.capitalize()}: {done}/{len(items)}")

        return target

    def _filter_types(
        self,
        data: Dict[str, Dict[str, Any]],
        include_dai_hoc: bool,
        include_cao_dang: bool,
    ) -> Dict[str, Dict[str, Any]]:
        result = {}
        for code, info in data.items():
            t = info.get("type", TYPE_DAI_HOC)
            if t == TYPE_CAO_DANG and not include_cao_dang:
                continue
            if t != TYPE_CAO_DANG and not include_dai_hoc:
                # hoc_vien + dai_hoc gộp nhóm đại học
                continue
            result[code] = info
        return result

    def _save_cache(self, data: Dict[str, Dict[str, Any]]):
        try:
            # Lưu toàn bộ (không filter) để lần sau dùng lại
            with open(self.cache_file, "w", encoding="utf-8") as f:
                json.dump(data, f, ensure_ascii=False, indent=2)
            print(f"[DANH_BA] Đã lưu cache: {self.cache_file}")
        except Exception as e:
            print(f"[CẢNH BÁO] Không lưu được cache danh bạ: {e}")

    def list_schools(
        self,
        school_type: Optional[str] = None,
        keyword: str = "",
    ) -> List[Dict[str, Any]]:
        """
        Trả về danh sách trường đã sắp xếp theo mã.
        school_type: None | 'dai_hoc' | 'cao_dang' | 'hoc_vien' | 'dai_hoc_hoc_vien'
        """
        if not self.schools:
            self.load()

        rows = []
        kw = keyword.strip().lower()
        for code, info in self.schools.items():
            t = info.get("type", TYPE_DAI_HOC)
            if school_type == "dai_hoc_hoc_vien" and t == TYPE_CAO_DANG:
                continue
            if school_type in (TYPE_DAI_HOC, TYPE_CAO_DANG, TYPE_HOC_VIEN) and t != school_type:
                continue
            if kw and kw not in code.lower() and kw not in (info.get("name") or "").lower():
                continue
            regions = split_addresses_by_region(info.get("dia_chi", ""))
            rows.append({
                "code": code,
                "name": info.get("name", ""),
                "slug": info.get("slug", ""),
                "type": t,
                "type_label": info.get("type_label") or TYPE_LABELS.get(t, t),
                "loai_truong": info.get("loai_truong") or classify_school_sector(info.get("name", "")),
                "thong_tin_chung": info.get("thong_tin_chung", ""),
                "dia_chi": info.get("dia_chi", ""),
                "khu_vuc": info.get("khu_vuc") or regions.get("khu_vuc", ""),
                "website": info.get("website", ""),
                "hotline": info.get("hotline", ""),
                "fanpage": info.get("fanpage", ""),
                "linh_vuc_chuong_trinh": info.get("linh_vuc_chuong_trinh", ""),
                "vi_the_thanh_tuu": info.get("vi_the_thanh_tuu", ""),
                "gioi_thieu_url": info.get("gioi_thieu_url", ""),
                "domain_diem_chuan": info.get("domain_diem_chuan", ""),
                "link_quy_che": info.get("link_quy_che", ""),
                "notices_checked": bool(info.get("notices_checked")),
            })

        rows.sort(key=lambda x: (x["type_label"], x["code"]))
        return rows

    def get_codes(
        self,
        school_type: Optional[str] = None,
    ) -> List[str]:
        """Danh sách mã trường (dùng làm đầu vào crawler)."""
        return [r["code"] for r in self.list_schools(school_type=school_type)]

    def export_json(
        self,
        output_path: str = "data/output/danh_sach_ma_truong.json",
        school_type: Optional[str] = None,
        also_update_config: bool = True,
    ) -> str:
        """Xuất JSON danh sách trường (+ tùy chọn ghi config/schools_all.json)."""
        rows = self.list_schools(school_type=school_type)
        os.makedirs(os.path.dirname(output_path) or ".", exist_ok=True)

        payload = {
            "description": "Danh sách mã trường và website chính thức của trường",
            "total": len(rows),
            "schools": [
                {
                    "code": r["code"],
                    "name": r["name"],
                    "short_name": r["name"][:40],
                    "online_slug": r["slug"],
                    "type": r["type"],
                    "type_label": r["type_label"],
                    "loai_truong": r.get("loai_truong", ""),
                    "thong_tin_chung": r.get("thong_tin_chung", ""),
                    "dia_chi": r.get("dia_chi", ""),
                    "khu_vuc": r.get("khu_vuc", ""),
                    "website": r.get("website", ""),
                    "hotline": r.get("hotline", ""),
                    "fanpage": r.get("fanpage", ""),
                    "linh_vuc_chuong_trinh": r.get("linh_vuc_chuong_trinh", ""),
                    "vi_the_thanh_tuu": r.get("vi_the_thanh_tuu", ""),
                    "gioi_thieu_url": r.get("gioi_thieu_url", ""),
                    "domain_diem_chuan": r.get("domain_diem_chuan", ""),
                    "link_quy_che": r.get("link_quy_che", ""),
                    "notices_checked": bool(r.get("notices_checked")),
                }
                for r in rows
            ],
            "years": [2021, 2022, 2023, 2024, 2025, 2026],
        }
        with open(output_path, "w", encoding="utf-8") as f:
            json.dump(payload, f, ensure_ascii=False, indent=2)
        print(f"[DANH_BA] Đã xuất JSON: {output_path} ({len(rows)} trường)")

        if also_update_config:
            config_path = "config/schools_all.json"
            os.makedirs(os.path.dirname(config_path), exist_ok=True)
            with open(config_path, "w", encoding="utf-8") as f:
                json.dump(payload, f, ensure_ascii=False, indent=2)
            print(f"[DANH_BA] Đã cập nhật config: {config_path}")

        return output_path

    def export_excel(
        self,
        output_path: str = "data/output/danh_sach_ma_truong.xlsx",
        school_type: Optional[str] = None,
    ) -> str:
        """Xuất Excel danh sách mã trường (dễ lọc / chọn để thu thập)."""
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter

        rows = self.list_schools(school_type=school_type)
        os.makedirs(os.path.dirname(output_path) or ".", exist_ok=True)

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Danh_Sach_Ma_Truong"

        headers = [
            "STT",
            "Mã trường",
            "Tên trường",
            "Loại hình",
            "Loại trường",
            "Slug (URL)",
            "Website chính thức",
            "Thông tin chung",
            "Địa chỉ các cơ sở",
            "Khu vực",
            "Website",
            "Hotline",
            "Fanpage",
            "Lĩnh vực & Chương trình đào tạo",
            "Vị thế & Thành tựu nổi bật",
            "Link giới thiệu",
            "Domain thông báo điểm chuẩn",
            "Link thông báo quy chế tuyển sinh",
        ]
        header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
        header_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        data_font = Font(name="Arial", size=10)
        thin = Border(
            left=Side(style="thin", color="D3D3D3"),
            right=Side(style="thin", color="D3D3D3"),
            top=Side(style="thin", color="D3D3D3"),
            bottom=Side(style="thin", color="D3D3D3"),
        )
        type_fills = {
            TYPE_DAI_HOC: PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid"),
            TYPE_HOC_VIEN: PatternFill(start_color="DDEBF7", end_color="DDEBF7", fill_type="solid"),
            TYPE_CAO_DANG: PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid"),
        }

        for col, title in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col, value=title)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        ws.row_dimensions[1].height = 32

        for idx, r in enumerate(rows, start=1):
            link = r.get("website") or ""
            vals = [
                idx,
                r["code"],
                r["name"],
                r["type_label"],
                r.get("loai_truong") or "",
                r["slug"],
                link,
                r.get("thong_tin_chung") or "",
                r.get("dia_chi") or "",
                r.get("khu_vuc") or "",
                r.get("website") or "",
                r.get("hotline") or "",
                r.get("fanpage") or "",
                r.get("linh_vuc_chuong_trinh") or "",
                r.get("vi_the_thanh_tuu") or "",
                r.get("gioi_thieu_url") or "",
                r.get("domain_diem_chuan") or "",
                r.get("link_quy_che") or "",
            ]
            fill = type_fills.get(r["type"])
            for c, val in enumerate(vals, start=1):
                cell = ws.cell(row=idx + 1, column=c, value=val)
                cell.font = data_font
                cell.border = thin
                cell.alignment = Alignment(
                    horizontal="center" if c in (1, 2, 4) else "left",
                    vertical="top",
                    wrap_text=c >= 7,
                )
                if fill and c <= 4:
                    cell.fill = fill
            # Hàng cao hơn nếu có nội dung dài
            if any(vals[i] for i in (6, 11, 12)):
                ws.row_dimensions[idx + 1].height = 60

        # Sheet thống kê
        ws2 = wb.create_sheet("Thong_Ke")
        stats = {}
        for r in rows:
            stats[r["type_label"]] = stats.get(r["type_label"], 0) + 1
        ws2["A1"] = "Loại hình"
        ws2["B1"] = "Số lượng"
        ws2["A1"].font = header_font
        ws2["B1"].font = header_font
        ws2["A1"].fill = header_fill
        ws2["B1"].fill = header_fill
        row_i = 2
        for label, cnt in sorted(stats.items()):
            ws2.cell(row=row_i, column=1, value=label).font = data_font
            ws2.cell(row=row_i, column=2, value=cnt).font = data_font
            row_i += 1
        ws2.cell(row=row_i, column=1, value="TỔNG").font = Font(name="Arial", bold=True)
        ws2.cell(row=row_i, column=2, value=len(rows)).font = Font(name="Arial", bold=True)
        with_profile = sum(1 for r in rows if r.get("thong_tin_chung") or r.get("website"))
        ws2.cell(row=row_i + 2, column=1, value="Có hồ sơ giới thiệu").font = data_font
        ws2.cell(row=row_i + 2, column=2, value=with_profile).font = data_font
        ws2.column_dimensions["A"].width = 28
        ws2.column_dimensions["B"].width = 12

        ws.freeze_panes = "A2"
        last_col = get_column_letter(len(headers))
        ws.auto_filter.ref = f"A1:{last_col}{len(rows) + 1}"
        widths = [6, 12, 40, 12, 40, 45, 40, 28, 28, 16, 28, 45, 45, 40, 36, 55]
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w

        wb.save(output_path)
        print(f"[DANH_BA] Đã xuất Excel: {output_path} ({len(rows)} trường)")
        return output_path

    def print_summary(self, school_type: Optional[str] = None, preview: int = 15):
        """In tóm tắt danh bạ ra console."""
        rows = self.list_schools(school_type=school_type)
        stats: Dict[str, int] = {}
        for r in rows:
            stats[r["type_label"]] = stats.get(r["type_label"], 0) + 1

        print("\n" + "=" * 70)
        print(" DANH BẠ MÃ TRƯỜNG ĐẠI HỌC / CAO ĐẲNG / HỌC VIỆN")
        print("=" * 70)
        for label, cnt in sorted(stats.items()):
            print(f"  - {label}: {cnt}")
        print(f"  → TỔNG: {len(rows)} trường")
        print("-" * 70)
        print(f"  Xem trước {min(preview, len(rows))} mã đầu:")
        for r in rows[:preview]:
            print(f"    {r['code']:<10} | {r['type_label']:<10} | {r['name'][:45]}")
        if len(rows) > preview:
            print(f"    ... và {len(rows) - preview} trường khác")
        print("=" * 70 + "\n")
